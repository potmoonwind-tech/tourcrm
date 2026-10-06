"""Stage 1 of the Sachiku move: everything the old program knows → export.json.

Runs in the old program's environment, on the copy: its own database (filled by run_sync.py from
the sheets) plus the raw BALANCE tabs read line by line with the old program's own rules
(categories, name fixes, which number is USD, what the colours mean)."""
import datetime as dt
import glob
import json
import os
import re
import sqlite3
import sys
from collections import Counter
from pathlib import Path

HERE = Path(__file__).parent
APP = HERE / "app"
os.chdir(APP)
sys.path.insert(0, str(APP))
import requests  # noqa: E402,F811
import hashlib  # noqa: E402

RAW = HERE / "raw"


def cached_get(url, *a, **k):          # the sheets were downloaded once by run_sync.py
    f = RAW / hashlib.sha1(url.encode()).hexdigest()[:16]
    r = requests.models.Response()
    r.status_code, r._content, r.url = (200, f.read_bytes(), url) if f.exists() else (404, b"", url)
    return r


requests.get = cached_get

from openpyxl import load_workbook  # noqa: E402

import archive_sync  # noqa: E402
import menu_data  # noqa: E402
import seed_data  # noqa: E402
from cancel_sync import fetch_cancelled_tour_codes  # noqa: E402
from debts_sync import PHASE_RE, _OK_RE, _shade  # noqa: E402
from menu_sync import fetch_menu  # noqa: E402
from profit_sync import (SHEET_IDS, TOUR_CODE_RE, _SUMMARY_KEYS, _categorize, _extract_pax, _fix_name, _norm_code,  # noqa: E402
                         _num, _usd_from_row)

out = {"made_at": dt.datetime.now().isoformat(timespec="seconds"), "notes": []}
note = out["notes"].append

# --- old database -------------------------------------------------------------------------
c = sqlite3.connect("gtc360.db")
c.row_factory = sqlite3.Row
rows = lambda q: [dict(r) for r in c.execute(q)]  # noqa: E731
out["contacts"] = {t: rows(f"select * from {t}") for t in
                   ("contacts_guides", "contacts_hotels", "contacts_restaurants", "contacts_stay", "contacts_company", "contacts_extra")}
out["payment_terms"] = rows("select * from payment_terms")
out["tours_2026"] = rows("select * from tours order by bus_start")
days = rows("select * from daily_log order by tour_code, date")
for t in out["tours_2026"]:
    t["days"] = [d for d in days if d["tour_code"] == t["code"]]
out["old_profit"] = {r["tour_code"]: r for r in rows("select tour_code, pax, profit_usd, vat_usd, spent_usd, revenue_usd, first_date, last_date from tour_profit")}

# --- series (routes) and menus, from the old code ------------------------------------------
out["series"] = {k: {"name": v.get("name", k), "duration": v.get("duration"), "color": v.get("color"),
                     "nights": {str(n): x for n, x in (v.get("nights") or {}).items()}}
                 for k, v in seed_data.SERIES.items()}
out["series_start_offset"] = getattr(seed_data, "SERIES_START_OFFSET", {})
try:
    menu_data.sync_menu_data(fetch_menu())
except Exception as e:
    note(f"menu workbook not read: {e}")
out["menus"] = menu_data.RESTAURANT_MENUS
out["menu_extra"] = {"PORTION_TABLE": getattr(menu_data, "PORTION_TABLE", None), "DISH_RATIOS": {" / ".join(str(x) for x in k) if isinstance(k, tuple) else str(k): v for k, v in getattr(menu_data, "DISH_RATIOS", {}).items()}}

# --- cancelled -----------------------------------------------------------------------------
out["cancelled"] = sorted(fetch_cancelled_tour_codes())

# --- BALANCE tabs, line by line -------------------------------------------------------------
DATE_TEXT = re.compile(r"(\d{1,2})[./](\d{1,2})[./](\d{4})")


def cell_text(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return (v.date() if isinstance(v, dt.datetime) else v).isoformat()
    return "" if v is None else str(v).strip()


def parse_tab(ws, code_re, norm, season_of=None):
    tab = {"title": ws.title, "lines": [], "summary": {}, "header": []}
    m = code_re.search(ws.title or "")
    tab["code"] = norm(m.group(1)) if m else None
    tab["pax"] = _extract_pax(ws.title, code_re)
    a1 = str(ws["A1"].value or "")
    pm = PHASE_RE.search(a1)
    tab["phase"] = int(pm.group(1)) if pm else None
    dates = []
    for row in ws.iter_rows():
        cells = list(row) + [None] * max(0, 9 - len(row))
        vals = [cell_text(x.value) if x is not None else "" for x in cells[:9]]
        if not any(vals):
            continue
        if not tab["header"] and any(code_re.search(v) for v in vals[1:4] if v):
            tab["header"] = [v for v in vals if v]
        d = vals[0] if re.fullmatch(r"\d{4}-\d{2}-\d{2}", vals[0] or "") else None
        if d:
            dates.append(d)
        # summary block
        for i, v in enumerate(vals):
            for key, k in (("დღგ სავარაუდო", "vat"), ("დაიხარჯა", "spent"), ("ტურის ღირებ", "revenue"), ("მოგება", "profit"),
                           ("ჩარიცხული", "received"), ("ჩასარიცხია", "to_receive")):
                if v.startswith(key):
                    nums = [_num(x) for x in vals[i + 1:] if _num(x) is not None]
                    if nums and k not in tab["summary"]:
                        tab["summary"][k] = nums[0]
                        if k == "profit" and len(nums) >= 3:
                            tab["summary"]["profit_after_vat"] = nums[2]
        name = _fix_name(vals[1])
        if not name or any(name.startswith(k) for k in _SUMMARY_KEYS) or code_re.search(name):
            continue
        rate, gel, usd_e, vat = _num(vals[2]), _num(vals[3]), _num(vals[4]), _num(vals[6])
        usd = _usd_from_row(vals)
        if not (usd or gel):
            # an itinerary row with no money (Armenia meals, "on their own"): keep for the days
            tab["lines"].append({"name": name, "date": d, "cat": _categorize(name), "info": True})
            continue
        e_shade, f_shade = _shade(cells[4]) if cells[4] is not None else "none", _shade(cells[5]) if cells[5] is not None else "none"
        status = vals[5]
        paid = f_shade == "green" and bool(_OK_RE.search(status))
        received = paid or e_shade == "green"
        if tab["phase"] in (2, 3):
            received = True
        if tab["phase"] == 3:
            paid = True
        h = vals[7]
        hd = DATE_TEXT.search(h or "")
        tab["lines"].append({
            "name": name, "date": d, "cat": _categorize(name), "rate": rate if rate and 2.0 < rate < 3.5 else None,
            "gel": gel, "usd": round(usd, 2) if usd else None, "usd_e": usd_e, "vat": vat,
            "received": received, "paid": paid, "status": status, "note": h,
            "paid_on": f"{hd.group(3)}-{int(hd.group(2)):02d}-{int(hd.group(1)):02d}" if hd else None,
            "e_shade": e_shade, "f_shade": f_shade,
        })
    run = archive_sync._run_of_dates(dates) if dates else []
    tab["first_date"], tab["last_date"] = (run[0], run[-1]) if run else (None, None)
    if season_of and run:
        tab["season"] = season_of(dt.date.fromisoformat(run[0]))
    return tab


def workbook(sheet_id):
    f = glob.glob(str(HERE / "xl" / f"{sheet_id}.xlsx"))
    return load_workbook(f[0], data_only=True) if f else None


out["tabs"] = []
for key, sid in SHEET_IDS.items():
    wb = workbook(sid)
    if wb is None:
        note(f"balance workbook {key} not downloaded")
        continue
    for ws in wb.worksheets:
        t = parse_tab(ws, TOUR_CODE_RE, _norm_code, archive_sync._season_year)
        if t["code"]:
            t["workbook"], t["kind"] = key, "current"
            out["tabs"].append(t)
        else:
            note(f"{key}: tab '{ws.title}' has no tour code — skipped")
for key, sid in archive_sync.ARCHIVE_SHEET_IDS.items():
    wb = workbook(sid)
    if wb is None:
        note(f"archive workbook {key} not downloaded")
        continue
    for ws in wb.worksheets:
        t = parse_tab(ws, archive_sync.ARCHIVE_CODE_RE, archive_sync._norm_archive_code, archive_sync._season_year)
        if not t["code"]:
            name = re.sub(r"\d+\s*\+\s*\d+", "", ws.title or "").strip()
            if not name:
                continue
            t["code"] = name                     # a one-off charter titled by the client's name
            t["charter"] = True
        t["workbook"], t["kind"] = key, "archive"
        out["tabs"].append(t)

Path(HERE / "export.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
cnt = Counter(t["kind"] for t in out["tabs"])
print("tabs:", dict(cnt), "lines:", sum(len(t["lines"]) for t in out["tabs"]),
      "tours_2026:", len(out["tours_2026"]), "series:", len(out["series"]), "menus:", len(out["menus"]),
      "cancelled:", len(out["cancelled"]), "notes:", len(out["notes"]))
for n in out["notes"][:20]:
    print(" -", n)
