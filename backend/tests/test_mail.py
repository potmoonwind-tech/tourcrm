"""Email from the CRM: your mailbox (app password, kept encrypted), templates, a booking's email filled in,
sending (faked SMTP here), the log, and the booking becoming "requested"."""
import smtplib

import pytest

from app import mail
from .test_tours import agency, mk

SENT: list = []


class FakeSMTP:
    def __init__(self, host, port, timeout=None, context=None):
        self.host, self.port = host, port

    def starttls(self, context=None):
        pass

    def login(self, user, password):
        if password != "abcdabcdabcdabcd":
            raise smtplib.SMTPAuthenticationError(535, b"bad")
        self.user = user

    def send_message(self, msg):
        SENT.append(msg)

    def quit(self):
        pass


@pytest.fixture(autouse=True)
def fake_smtp(monkeypatch):
    SENT.clear()
    monkeypatch.setattr(mail.smtplib, "SMTP", FakeSMTP)
    monkeypatch.setattr(mail.smtplib, "SMTP_SSL", FakeSMTP)


def setup_booking(c, base):
    hotel = mk(c, f"{base}/vendors", {"name": "Hotel Rustaveli", "type": "hotel", "email": "res@rustaveli.ge", "contact_person": "Nana"})
    route = mk(c, f"{base}/series", {"code": "R", "days": [
        {"city": "Tbilisi", "hotel_id": hotel["id"]}, {"city": "Tbilisi", "hotel_id": hotel["id"]}, {"city": "Tbilisi"}]})
    t = mk(c, f"{base}/tours", {"code": "ZK-1011", "series_id": route["id"], "start_date": "2026-10-11", "pax": 20, "rooms": "9 DBL + 2 SGL"})
    b = next(x for x in c.get(f"{base}/tours/{t['id']}/bookings").json()["bookings"] if x["vendor"]["name"] == "Hotel Rustaveli")
    return t, b


def test_connect_draft_send_and_log(make_client):
    c, base, _ = agency(make_client)
    assert c.get(f"{base}/mail/account").json()["connected"] is False
    bad = c.put(f"{base}/mail/account", json={"provider": "gmail", "email": "nino@gmail.com", "password": "wrong"})
    assert bad.status_code == 422 and "app password" in bad.json()["detail"]
    ok = c.put(f"{base}/mail/account", json={"provider": "gmail", "email": "nino@gmail.com", "password": "abcd abcd abcd abcd",
                                            "from_name": "Nino · GT360", "signature": "Nino Beridze\nGT360 Travel"}).json()
    assert ok["connected"] and ok["email"] == "nino@gmail.com" and "password" not in ok and ok["last_ok_at"]

    t, b = setup_booking(c, base)
    d = c.post(f"{base}/tours/{t['id']}/bookings/{b['id']}/email-draft", json={"language": "en"}).json()
    assert d["to"] == "res@rustaveli.ge" and d["kind"] == "booking_request" and d["has_account"]
    assert d["subject"] == "Booking request · ZK-1011 · 11 Oct 2026 – 13 Oct 2026"
    assert "Dear Nana" in d["body"] and "Check-in 11 Oct 2026, check-out 13 Oct 2026 · 2 nights" in d["body"]
    assert "Rooms: 9 DBL + 2 SGL" in d["body"] and d["body"].endswith("Nino Beridze\nGT360 Travel")
    ka = c.post(f"{base}/tours/{t['id']}/bookings/{b['id']}/email-draft", json={"language": "ka"}).json()
    assert ka["subject"].startswith("ჯავშნის მოთხოვნა") and "ოთახები: 9 DBL + 2 SGL" in ka["body"]

    r = c.post(f"{base}/mail/send", json={"to": [d["to"]], "subject": d["subject"], "body": d["body"],
                                         "tour_id": t["id"], "booking_id": b["id"], "kind": d["kind"]})
    assert r.status_code == 200, r.text
    msg = SENT[-1]
    assert msg["To"] == "res@rustaveli.ge" and msg["Reply-To"] == "nino@gmail.com" and "Nino" in msg["From"]
    assert "9 DBL + 2 SGL" in msg.get_content()
    after = next(x for x in c.get(f"{base}/tours/{t['id']}/bookings").json()["bookings"] if x["id"] == b["id"])
    assert after["status"] == "requested"
    log = c.get(f"{base}/mail/log?tour_id={t['id']}").json()
    assert len(log) == 1 and log[0]["error"] == "" and log[0]["vendor_id"] == b["vendor"]["id"]


def test_templates_are_the_agencys_own_and_go_back(make_client):
    c, base, _ = agency(make_client)
    rows = c.get(f"{base}/mail/templates").json()
    assert len(rows) == len(mail.KINDS) * len(mail.LANGS) and not any(r["custom"] for r in rows)
    r = c.put(f"{base}/mail/templates/booking_request/en", json={"subject": "Rooms for {group}", "body": "Hi {supplier}, {what}"})
    assert r.json()["custom"]
    c.put(f"{base}/mail/account", json={"provider": "gmail", "email": "n@gmail.com", "password": "abcdabcdabcdabcd"})
    t, b = setup_booking(c, base)
    d = c.post(f"{base}/tours/{t['id']}/bookings/{b['id']}/email-draft", json={"language": "en"}).json()
    assert d["subject"] == "Rooms for ZK-1011" and d["body"].startswith("Hi Nana, • Check-in")
    assert c.delete(f"{base}/mail/templates/booking_request/en").status_code == 204
    assert not next(x for x in c.get(f"{base}/mail/templates").json() if x["kind"] == "booking_request" and x["language"] == "en")["custom"]
    assert c.put(f"{base}/mail/templates/nonsense/en", json={"subject": "x", "body": "y"}).status_code == 404


def test_no_mailbox_no_sending_and_failures_are_logged(make_client, monkeypatch):
    c, base, _ = agency(make_client)
    assert c.post(f"{base}/mail/send", json={"to": ["a@b.ge"], "subject": "s", "body": "b"}).status_code == 409
    c.put(f"{base}/mail/account", json={"provider": "gmail", "email": "n@gmail.com", "password": "abcdabcdabcdabcd"})

    def refuse(self, msg):
        raise smtplib.SMTPRecipientsRefused({"x@y.ge": (550, b"no")})
    monkeypatch.setattr(FakeSMTP, "send_message", refuse)
    r = c.post(f"{base}/mail/send", json={"to": ["x@y.ge"], "subject": "s", "body": "b"})
    assert r.status_code == 422 and "x@y.ge" in r.json()["detail"]
    assert c.get(f"{base}/mail/log").json()[0]["error"]
    # Another server: never an internal address.
    r = c.put(f"{base}/mail/account", json={"provider": "other", "email": "n@x.ge", "password": "p", "smtp_host": "localhost", "smtp_port": 587})
    assert r.status_code == 422


def test_the_rooming_list_goes_to_the_hotel(make_client):
    """Rooms in order, the bed in words, the note for the hotel - attached to the hotel's email."""
    import io
    from openpyxl import load_workbook
    c, base, _ = agency(make_client)
    c.put(f"{base}/mail/account", json={"provider": "gmail", "email": "nino@gmail.com", "password": "abcd abcd abcd abcd"})
    t, b = setup_booking(c, base)
    url = f"{base}/tours/{t['id']}/participants"
    for name, room, kind, note in [("Ana Berg", "R10", "TWN", ""), ("Leo Berg", "R10", "TWN", ""),
                                   ("Mia Holm", "R2", "DBL", "Couple: one big bed"), ("Tom Holm", "R2", "DBL", "Couple: one big bed")]:
        mk(c, url, {"full_name": name, "room_group": room, "room_type": kind, "room_note": note})
    assert c.get(url).json()[2]["room_note"] == "Couple: one big bed"

    rows = list(load_workbook(io.BytesIO(c.get(f"{base}/tours/{t['id']}/rooming-list.xlsx").content)).active.iter_rows(min_row=4, values_only=True))
    assert [(r[1], r[2], r[3], r[4], r[10]) for r in rows] == [
        ("2", "DBL", "1 double bed", "Mia Holm", "Couple: one big bed"), ("2", "DBL", "1 double bed", "Tom Holm", "Couple: one big bed"),
        ("10", "TWN", "2 single beds", "Ana Berg", None), ("10", "TWN", "2 single beds", "Leo Berg", None)]

    d = c.post(f"{base}/tours/{t['id']}/bookings/{b['id']}/email-draft", json={"language": "en"}).json()
    assert d["can_attach_rooming"] is True
    r = c.post(f"{base}/mail/send", json={"to": [d["to"]], "subject": d["subject"], "body": d["body"],
                                         "tour_id": t["id"], "booking_id": b["id"], "attach_rooming": True})
    assert r.status_code == 200, r.text
    [att] = list(SENT[-1].iter_attachments())
    assert att.get_filename() == "ZK-1011-rooming-list.xlsx" and len(att.get_content()) > 1000


def test_one_email_for_the_whole_series(make_client):
    from datetime import date, timedelta
    c, base, _ = agency(make_client)
    hotel = mk(c, f"{base}/vendors", {"name": "Rooms Hotel", "type": "hotel", "email": "res@rooms.ge", "contact_person": "Nino"})
    route = mk(c, f"{base}/series", {"code": "HM", "name": "Highlights", "days": [{"city": "Tbilisi", "hotel_id": hotel["id"]}, {"city": "Tbilisi"}]})
    soon = date.today() + timedelta(days=40)
    tours = [mk(c, f"{base}/tours", {"series_id": route["id"], "start_date": (soon + timedelta(days=7 * i)).isoformat(), "pax": 10 + i, "rooms": "5 DBL"})
             for i in range(3)]
    d = c.post(f"{base}/series/{route['id']}/bookings/{hotel['id']}/email-draft", json={"language": "en"}).json()
    assert d["to"] == "res@rooms.ge" and len(d["tours"]) == 3 and all(x["included"] for x in d["tours"])
    assert d["subject"].startswith("Booking request · series HM (Highlights) · 3 groups")
    # Each departure: its number, the tour's code, dates and people.
    assert "1. Tour " + tours[0]["code"] in d["body"] and "11 people" in d["body"] and "3. Tour " + tours[2]["code"] in d["body"]
    assert "Check-in" in d["body"] and "Dear Nino" in d["body"]
    # Only two of them.
    two = c.post(f"{base}/series/{route['id']}/bookings/{hotel['id']}/email-draft",
                 json={"language": "en", "tour_ids": [tours[0]["id"], tours[2]["id"]]}).json()
    assert "2 groups" in two["subject"] and tours[1]["code"] not in two["body"]
    # Sent once; every departure's booking becomes requested and shows the email.
    c.put(f"{base}/mail/account", json={"provider": "gmail", "email": "nino@gmail.com", "password": "abcd abcd abcd abcd"})
    ids = [x["booking_id"] for x in d["tours"]]
    r = c.post(f"{base}/mail/send", json={"to": ["res@rooms.ge"], "subject": d["subject"], "body": d["body"],
                                          "kind": "series_request", "booking_ids": ids})
    assert r.status_code == 200, r.text
    assert len(SENT) == 1
    for x in tours:
        b = next(b for b in c.get(f"{base}/tours/{x['id']}/bookings").json()["bookings"] if b["vendor"]["id"] == hotel["id"])
        assert b["status"] == "requested"
        assert len(c.get(f"{base}/mail/log?tour_id={x['id']}").json()) == 1
