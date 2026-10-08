import io
from datetime import date

from openpyxl import Workbook, load_workbook

from .conftest import invite_and_join, signup
from .test_tours import agency, mk, seed_route


def tour_with_route(c, base, start="2026-09-06"):
    *_, series = seed_route(c, base)  # 3-day route
    return mk(c, f"{base}/tours", {"series_id": series["id"], "start_date": start})


# --- staff ---------------------------------------------------------------------------

def test_staff_by_day_range_and_day_view(make_client):
    c, base, _ = agency(make_client)
    t = tour_with_route(c, base)
    g1 = mk(c, f"{base}/vendors", {"name": "Wang", "type": "guide"})
    g2 = mk(c, f"{base}/vendors", {"name": "Levan", "type": "guide"})
    bus = mk(c, f"{base}/vendors", {"name": "Coaches", "type": "transport"})
    mk(c, f"{base}/tours/{t['id']}/staff", {"vendor_id": g1["id"], "role": "guide", "start_day": 1, "end_day": 2})
    mk(c, f"{base}/tours/{t['id']}/staff", {"vendor_id": g2["id"], "role": "guide", "start_day": 3})
    mk(c, f"{base}/tours/{t['id']}/staff", {"vendor_id": bus["id"], "role": "transport"})

    def today_staff(day):
        [row] = c.get(f"{base}/day/{day}").json()
        return sorted(s["vendor"]["name"] for s in row["staff_today"])

    assert today_staff("2026-09-07") == ["Coaches", "Wang"]
    assert today_staff("2026-09-08") == ["Coaches", "Levan"]

    # Ranges are day numbers, so they follow the tour when it moves.
    c.patch(f"{base}/tours/{t['id']}", json={"start_date": "2026-10-01"})
    assert today_staff("2026-10-03") == ["Coaches", "Levan"]

    # Out-of-range and reversed ranges are refused.
    bad = c.post(f"{base}/tours/{t['id']}/staff", json={"vendor_id": g1["id"], "role": "guide", "start_day": 4})
    assert bad.status_code == 422
    bad = c.post(f"{base}/tours/{t['id']}/staff", json={"vendor_id": g1["id"], "role": "guide", "start_day": 2, "end_day": 1})
    assert bad.status_code == 422


def test_shortening_tour_trims_staff(make_client):
    c, base, _ = agency(make_client)
    t = tour_with_route(c, base)
    g = mk(c, f"{base}/vendors", {"name": "G", "type": "guide"})
    late = mk(c, f"{base}/tours/{t['id']}/staff", {"vendor_id": g["id"], "role": "guide", "start_day": 3})
    span = mk(c, f"{base}/tours/{t['id']}/staff", {"vendor_id": g["id"], "role": "assistant", "start_day": 1, "end_day": 3})
    t2 = c.patch(f"{base}/tours/{t['id']}", json={"end_date": "2026-09-07"}).json()
    ids = {s["id"]: s for s in t2["staff"]}
    assert late["id"] not in ids
    assert ids[span["id"]]["end_day"] is None  # now "until the end"


# --- participants & rooming ------------------------------------------------------------

def test_participants_and_rooming(make_client):
    c, base, _ = agency(make_client)
    t = tour_with_route(c, base)
    url = f"{base}/tours/{t['id']}/participants"
    for name, rt, grp in [("A", "DBL", "R1"), ("B", "DBL", "R1"), ("C", "DBL", ""), ("D", "DBL", ""),
                          ("E", "SGL", ""), ("F", "", "R9"), ("G", "", "")]:
        mk(c, url, {"full_name": name, "room_type": rt, "room_group": grp})
    mk(c, url, {"full_name": "H", "passport_expiry": "2026-12-01"})  # < tour end + 6 months

    r = c.get(f"{base}/tours/{t['id']}/rooming").json()
    # R1 = 1 DBL, C+D packed into 1 DBL, E = SGL, R9 alone = SGL; G and H unassigned
    assert r["rooms"] == {"DBL": 2, "SGL": 2}
    assert r["text"] == "2 DBL + 2 SGL"
    assert r["unassigned"] == 2
    assert r["passport_warnings"] == 1
    assert c.get(f"{base}/tours/{t['id']}").json()["participant_count"] == 8
    assert next(x for x in c.get(f"{base}/tours").json() if x["id"] == t["id"])["participant_count"] == 8

    people = c.get(url).json()
    assert [p["full_name"] for p in people] == list("ABCDEFGH")
    upd = c.patch(f"{url}/{people[0]['id']}", json={"passport_number": "E1", "birth_date": "1990-01-02"}).json()
    assert upd["passport_number"] == "E1" and upd["birth_date"] == "1990-01-02"
    upd = c.patch(f"{url}/{people[0]['id']}", json={"birth_date": None}).json()
    assert upd["birth_date"] is None

    x = c.get(f"{base}/tours/{t['id']}/rooming-list.xlsx")
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["E4"].value == "A" and ws["B4"].value == "1"  # grouped rooms first ("R1" reads as room 1)

    assert c.delete(url).status_code == 204
    assert c.get(url).json() == []


def test_participants_are_scoped_and_hidden_from_guides(make_client):
    a, base_a, ws_a = agency(make_client, "a@a.com", "A")
    b, base_b, _ = agency(make_client, "b@b.com", "B")
    t = tour_with_route(a, base_a)
    p = mk(a, f"{base_a}/tours/{t['id']}/participants", {"full_name": "Secret", "passport_number": "X1"})
    assert b.get(f"{base_b}/tours/{t['id']}/participants").status_code == 404
    assert b.patch(f"{base_b}/tours/{t['id']}/participants/{p['id']}", json={"full_name": "x"}).status_code == 404
    assert b.get(f"{base_b}/tours/{t['id']}/rooming-list.xlsx").status_code == 404

    guide = make_client()
    signup(guide, "g@a.com")
    invite_and_join(a, ws_a, guide, "g@a.com", "guide")
    assert guide.get(f"{base_a}/tours/{t['id']}/participants").status_code == 403


# --- import ------------------------------------------------------------------------------

def xlsx(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(c, url, data, name="list.xlsx", **params):
    return c.post(url, params=params, files={"file": (name, data)})


def test_participant_import_preview_then_commit(make_client):
    c, base, _ = agency(make_client)
    t = tour_with_route(c, base)
    data = xlsx([
        ["Group HM-0906 passenger list"],             # title row above the header
        ["No.", "姓名", "性别", "出生日期", "护照号", "Passport expiry", "房型", "Room no"],
        [1, "WANG Li", "女", date(1985, 4, 12), "E123", "2031-01-31", "双人间", "R1"],
        [2, "ZHANG Wei", "男", "12.05.1980", "E124", "31/01/2031", "DBL", "R1"],
        [3, "LI Na", "F", "not a date", "E125", "", "SGL", ""],
        [4, "", "", "", "", "", "", ""],               # empty -> ignored
        [5, "CHEN Bo", "M", "", "E124", "06JAN2027", "Penthouse", ""],
    ])
    url = f"{base}/import/participants"
    r = upload(c, url, data, tour_id=t["id"]).json()
    assert r["committed"] is False and r["created"] == 0
    assert r["columns"]["姓名"] == "full_name" and r["columns"]["Room no"] == "room_group"
    assert r["columns"]["No."] == "no"          # the number column: where "领队" marks the tour leader
    assert r["total"] == 4 and r["valid"] == 2
    assert {e["row"] for e in r["errors"]} == {5, 7}  # bad date, unknown room type
    assert r["preview"][0]["gender"] == "F" and r["preview"][0]["room_type"] == "DBL"
    assert r["preview"][1]["birth_date"] == "1980-05-12"
    assert c.get(f"{base}/tours/{t['id']}/participants").json() == []  # preview saved nothing

    r = upload(c, url, data, tour_id=t["id"], commit="true").json()
    assert r["created"] == 2
    assert c.get(f"{base}/tours/{t['id']}/rooming").json()["text"] == "1 DBL"

    # replace mode swaps the list instead of appending
    small = xlsx([["full name", "room type"], ["SOLO Traveller", "SGL"]])
    upload(c, url, small, tour_id=t["id"], commit="true", mode="replace")
    assert [p["full_name"] for p in c.get(f"{base}/tours/{t['id']}/participants").json()] == ["SOLO Traveller"]


def test_first_and_last_name_columns_and_csv(make_client):
    c, base, _ = agency(make_client)
    t = tour_with_route(c, base)
    csv_data = "Surname;Given name;Passport\nkim;Ji-ho;M1\n".encode("utf-8-sig")
    r = upload(c, f"{base}/import/participants", csv_data, "list.csv", tour_id=t["id"], commit="true").json()
    assert r["created"] == 1
    assert c.get(f"{base}/tours/{t['id']}/participants").json()[0]["full_name"] == "KIM Ji-ho"


def test_tour_import(make_client):
    c, base, _ = agency(make_client)
    *_, series = seed_route(c, base)
    mk(c, f"{base}/clients", {"name": "Orient Holidays"})
    data = xlsx([
        ["Route", "Start", "Client", "Pax", "Status"],
        ["HM", date(2026, 11, 6), "orient holidays", 24, "confirmed"],
        ["HM", date(2026, 11, 6), "Brand New Agency", 10, "option"],
        ["XX", date(2026, 11, 7), "", 5, ""],
        ["", date(2026, 11, 8), "", 5, ""],
    ])
    r = upload(c, f"{base}/import/tours", data, commit="true").json()
    assert r["created"] == 2
    assert {e["row"] for e in r["errors"]} == {4, 5}  # unknown route; no route and no end date
    assert any("Brand New Agency" in w["message"] for w in r["warnings"])
    tours = {t["code"]: t for t in c.get(f"{base}/tours").json()}
    assert set(tours) == {"HM-1106", "HM-1106B"}
    assert tours["HM-1106"]["client_name"] == "Orient Holidays"  # matched case-insensitively
    assert tours["HM-1106B"]["status"] == "inquiry"
    assert len(c.get(f"{base}/tours/{tours['HM-1106']['id']}").json()["days"]) == 3
    # importing the same file again: codes now taken
    r = upload(c, f"{base}/import/tours", xlsx([["code", "start", "end"], ["HM-1106", "2026-11-06", "2026-11-07"]])).json()
    assert r["valid"] == 0 and "already exists" in r["errors"][0]["message"]


def test_vendor_and_client_import_permissions_and_duplicates(make_client):
    owner, base, ws = agency(make_client)
    staff = make_client()
    signup(staff, "s@a.com")
    invite_and_join(owner, ws, staff, "s@a.com", "staff")
    data = xlsx([["Name", "Type", "City", "Payment days"], ["Hotel X", "სასტუმრო", "Tbilisi", 7],
                 ["Hotel X", "hotel", "", ""], ["Bus Y", "bus", "", "x"]])
    assert upload(staff, f"{base}/import/vendors", data).status_code == 403
    r = upload(owner, f"{base}/import/vendors", data, commit="true").json()
    assert r["created"] == 1
    assert {e["row"] for e in r["errors"]} == {3, 4}  # duplicate in file; bad number
    assert owner.get(f"{base}/vendors").json()[0]["payment_days"] == 7

    r = upload(owner, f"{base}/import/clients", xlsx([["name", "country"], ["Seoul TP", "Korea"]]), commit="true").json()
    assert r["created"] == 1


def test_import_rejects_bad_files_and_missing_columns(make_client):
    c, base, _ = agency(make_client)
    t = tour_with_route(c, base)
    assert upload(c, f"{base}/import/clients", b"not excel", "x.xlsx").status_code == 422
    assert upload(c, f"{base}/import/clients", b"a,b", "x.pdf").status_code == 422
    r = upload(c, f"{base}/import/participants", xlsx([["foo", "bar"], [1, 2]]), tour_id=t["id"]).json()
    assert r["valid"] == 0 and "Missing column" in r["errors"][0]["message"]
    assert upload(c, f"{base}/import/participants", xlsx([["name"], ["x"]])).status_code == 422  # no tour_id


def test_import_cannot_target_other_workspace_tour(make_client):
    a, base_a, _ = agency(make_client, "a@a.com", "A")
    b, base_b, _ = agency(make_client, "b@b.com", "B")
    t = tour_with_route(a, base_a)
    r = upload(b, f"{base_b}/import/participants", xlsx([["name"], ["Intruder"]]), tour_id=t["id"], commit="true")
    assert r.status_code == 404
    assert a.get(f"{base_a}/tours/{t['id']}/participants").json() == []


def test_templates_download(make_client):
    c, base, _ = agency(make_client)
    for entity in ("participants", "tours", "vendors", "clients"):
        r = c.get(f"{base}/import/{entity}/template")
        assert r.status_code == 200
        ws = load_workbook(io.BytesIO(r.content)).active
        assert ws["A1"].value  # header row present
    # The template round-trips through the importer.
    t = tour_with_route(c, base)
    tpl = c.get(f"{base}/import/participants/template").content
    r = upload(c, f"{base}/import/participants", tpl, tour_id=t["id"]).json()
    assert r["valid"] == 1, r


def test_a_partners_list_with_merged_rooms_a_leader_and_passports_on_sheet_two(make_client):
    """A list as partners send it: a title, "英文名" over SURNAME | Given, the tour leader marked "领队",
    roommates sharing one merged "标间/twin" cell, a note next to the room, passports on the second sheet."""
    c, base, _ = agency(make_client)
    t = tour_with_route(c, base)
    wb = Workbook()
    ws = wb.active
    for r in (["0928 group list"], ["序号", "中文名", "英文名", None, "性别", "出生日期", "联系方式", "分房", None],
              ["领队", "张三", "ZHANG", "SAN", "男", date(1985, 1, 2), "13800000000", "标间/twin", None],
              [1, "李四", "LI", "SI", "男", date(1970, 3, 4), None, None, None],
              [2, "王五", "WANG", "WU", "女", date(1960, 5, 6), None, "大床/double", "生日蛋糕"],
              [3, "赵六", "ZHAO", "LIU", "女", date(1962, 7, 8), None, None, None],
              [4, "钱七", "QIAN", "QI", "女", date(1990, 9, 10), None, "单间/single", None]):
        ws.append(r)
    ws.merge_cells("A1:I1")
    ws.merge_cells("C2:D2")
    ws.merge_cells("H3:H4")
    ws.merge_cells("H5:H6")
    p2 = wb.create_sheet("Sheet2")
    for r in (["张三", "护照", "EA1111111", date(1985, 1, 2), "男"], ["李四", "护照", "EB2222222", date(1970, 3, 4), "男"],
              ["王五", "护照", "EC3333333", date(1960, 5, 6), "女"], ["钱七", "护照", "ED4444444", date(1990, 9, 10), "女"]):
        p2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    r = upload(c, f"{base}/import/participants", buf.getvalue(), tour_id=t["id"], commit="true").json()
    assert r["created"] == 5 and r["errors"] == [], r
    people = {p["local_name"]: p for p in c.get(f"{base}/tours/{t['id']}/participants").json()}
    assert people["张三"]["full_name"] == "ZHANG SAN" and people["张三"]["notes"].startswith("Tour leader")
    assert people["张三"]["passport_number"] == "EA1111111" and people["赵六"]["passport_number"] == ""
    # Roommates: the merged cell, a twin for the leader and Li; a double for Wang and Zhao (its note too); Qian alone.
    assert people["张三"]["room_group"] == people["李四"]["room_group"] and people["李四"]["room_type"] == "TWN"
    assert people["王五"]["room_group"] == people["赵六"]["room_group"] and people["赵六"]["room_note"] == "生日蛋糕"
    assert people["钱七"]["room_type"] == "SGL" and people["钱七"]["room_group"] not in (people["张三"]["room_group"], people["王五"]["room_group"])
    assert c.get(f"{base}/tours/{t['id']}/rooming").json()["text"] == "1 DBL + 1 TWN + 1 SGL"
    assert c.get(f"{base}/tours/{t['id']}").json()["leaders"] == 1       # the free tour leader of the group


def test_a_list_nobody_taught_us_is_read_by_ai(make_client, monkeypatch):
    """Headers the usual reading doesn't know: AI says which column is which (it gets the first rows with the
    numbers masked), and the list is read as usual from that."""
    from app import importer
    seen = {}

    class Fake:
        def __init__(self):
            self.messages = self

        def create(self, model, max_tokens, messages):
            seen["prompt"] = messages[0]["content"]
            from types import SimpleNamespace
            text = '{"header_row": 1, "columns": {"A": "full_name", "B": "passport_number", "C": "room_type"}, "leader": "TL"}'
            return SimpleNamespace(content=[SimpleNamespace(type="text", text=text)])

    monkeypatch.setattr("app.programme.make_client", lambda ws: Fake())
    monkeypatch.setattr("app.ai.keys.key_for", lambda ws: "sk-test")
    c, base, _ = agency(make_client)
    t = tour_with_route(c, base)
    data = xlsx([["Persona", "Documento", "Habitación"], ["GARCIA Ana", "X1234567", "doble"], ["LOPEZ Juan", "X7654321", "doble"]])
    r = upload(c, f"{base}/import/participants", data, tour_id=t["id"]).json()
    assert r["by_ai"] is True and r["valid"] == 2 and r["columns"] == {"Persona": "full_name", "Documento": "passport_number", "Habitación": "room_type"}
    assert "X1234567" not in seen["prompt"] and "X9999999" in seen["prompt"]          # passport numbers never leave
    assert importer  # noqa
