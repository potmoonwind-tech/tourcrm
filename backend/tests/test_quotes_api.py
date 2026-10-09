import io

from openpyxl import load_workbook

from .conftest import invite_and_join, signup
from .test_finance import set_prices
from .test_tours import agency, mk


def new_quote(c, base, preset="hm"):
    return mk(c, f"{base}/quotes", {"preset": preset})


def doc(q):
    """The editable part of a QuoteOut."""
    return {k: v for k, v in q.items() if k not in ("id", "updated_at", "result")}


def test_create_from_preset_and_save(make_client):
    c, base, _ = agency(make_client)
    q = new_quote(c, base)
    assert q["name"] == "HM Georgia–Armenia 2026"
    assert q["result"]["errors"] == []
    by_n = {t["tourists"]: t for t in q["result"]["tiers"]}
    assert by_n[19]["label"] == "19+1" and by_n[19]["vehicle"] == "45-seater bus"

    d = doc(q)
    d["variables"]["armenia_meals"] = "0"
    d["name"] = "HM 2026 · no Armenian meals"
    saved = c.put(f"{base}/quotes/{q['id']}", json=d).json()
    assert saved["name"] == "HM 2026 · no Armenian meals"
    assert {t["tourists"]: t for t in saved["result"]["tiers"]}[19]["cost_dbl"] < by_n[19]["cost_dbl"]
    assert [x["name"] for x in c.get(f"{base}/quotes").json()] == ["HM 2026 · no Armenian meals"]

    copy = mk(c, f"{base}/quotes", {"copy_of": q["id"]})
    assert copy["name"].endswith("(copy)") and copy["variables"]["armenia_meals"] == "0"


def test_preview_and_breakdown_do_not_save(make_client):
    c, base, _ = agency(make_client)
    q = new_quote(c, base)
    d = doc(q)
    d["profit_min_pp"] = "100"
    r = c.post(f"{base}/quotes/preview", json=d).json()
    assert all(t["profit_pp"] >= 100 for t in r["tiers"])
    b = c.post(f"{base}/quotes/preview/breakdown", params={"tourists": 19}, json=d).json()
    assert {i["name"]: i["amount"] for i in b["items"]}["Guide"] == 1350
    assert c.post(f"{base}/quotes/preview/breakdown", params={"tourists": 40}, json=d).status_code == 422
    assert float(c.get(f"{base}/quotes/{q['id']}").json()["profit_min_pp"]) == 34  # unchanged on the server

    d["items"][0]["quantity"] = "days * nope"
    assert any("nope" in e for e in c.post(f"{base}/quotes/preview", json=d).json()["errors"])
    d["variables"]["days"] = "3"  # built-in names can't be redefined
    assert c.post(f"{base}/quotes/preview", json=d).status_code == 422


def test_validation_rejects_bad_documents(make_client):
    c, base, _ = agency(make_client)
    d = doc(new_quote(c, base))
    bad = dict(d, pax_min=30, pax_max=10)
    assert c.post(f"{base}/quotes/preview", json=bad).status_code == 422
    bad = dict(d, items=[dict(d["items"][0], basis="per_vehicle", capacity=None)])
    assert c.post(f"{base}/quotes/preview", json=bad).status_code == 422
    bad = dict(d, languages=["english!"])
    assert c.post(f"{base}/quotes/preview", json=bad).status_code == 422


def test_route_fill_uses_route_hotels_and_vendor_prices(make_client):
    c, base, _ = agency(make_client)
    hotel = mk(c, f"{base}/vendors", {"name": "Hualing", "type": "hotel"})
    set_prices(c, base, hotel["id"], room_DBL=("75", "USD"), room_SGL=("70", "USD"))
    resto = mk(c, f"{base}/vendors", {"name": "Supra", "type": "restaurant"})
    set_prices(c, base, resto["id"], dinner=("30", "GEL"))
    s = mk(c, f"{base}/series", {"code": "HM", "days": [
        {"city": "Tbilisi", "hotel_id": hotel["id"], "dinner_id": resto["id"]},
        {"city": "Tbilisi", "hotel_id": hotel["id"], "dinner_id": resto["id"]}, {"city": "Tbilisi"}]})
    r = c.post(f"{base}/quotes/route-fill", json={"series_id": s["id"], "meals": True}).json()
    assert r["days"] == 3 and len(r["nights"]) == 2
    assert float(r["nights"][0]["price_dbl"]) == 75 and float(r["nights"][0]["price_sgl"]) == 70
    dinner = next(i for i in r["items"] if i["name"] == "Dinner: Supra")
    assert dinner["quantity"] == "2" and dinner["currency"] == "GEL" and dinner["basis"] == "per_person"
    assert r["warnings"] == ["4 meals have no restaurant yet: they are in the calculation without a price"]


def test_price_list_json_and_excel(make_client):
    c, base, _ = agency(make_client)
    q = new_quote(c, base)
    zh = c.get(f"{base}/quotes/{q['id']}/price-list", params={"lang": "zh"}).json()
    assert zh["labels"]["col_dbl"] == "双人间/标间 每人" and zh["title"].startswith("格鲁吉亚")
    assert zh["labels"]["meals"] == "餐食" and zh["texts"]["meals"].startswith("含午餐")  # heading ≠ paragraph
    assert len(zh["tiers"]) == 17 and len(zh["hotels"]) == 8
    # An unknown language falls back to English labels and texts.
    fr = c.get(f"{base}/quotes/{q['id']}/price-list", params={"lang": "fr"}).json()
    assert fr["labels"]["col_hotel"] == "Hotel"
    x = c.get(f"{base}/quotes/{q['id']}/price-list.xlsx", params={"lang": "en"})
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    values = [v for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    assert "HM TOURS Georgia–Armenia 8 Nights" in values and "45-seater bus" in values


def test_create_tour_and_invoice_from_quote_then_budget(make_client):
    c, base, _ = agency(make_client)
    cl = mk(c, f"{base}/clients", {"name": "Orient"})
    q = new_quote(c, base)
    d = doc(q)
    d["client_id"] = cl["id"]
    c.put(f"{base}/quotes/{q['id']}", json=d)
    price = {t["tourists"]: t for t in q["result"]["tiers"]}[19]

    r = mk(c, f"{base}/quotes/{q['id']}/create-tour", {"tourists": 19, "start_date": "2026-11-01", "singles": 1})
    t = c.get(f"{base}/tours/{r['tour_id']}").json()
    assert (t["pax"], t["leaders"]) == (19, 1) and len(t["days"]) == 9 and t["days"][0]["city"] == "Tbilisi"   # 19+1
    assert t["rooms"] == "9 DBL + 2 SGL"
    inv = c.get(f"{base}/invoices/{r['invoice_id']}").json()
    assert inv["status"] == "draft" and inv["client"]["name"] == "Orient"
    assert inv["lines"][0]["quantity"] == 19 and inv["lines"][0]["unit_price"] == price["price_dbl"]
    assert inv["lines"][1]["unit_price"] == price["single_supplement"]
    assert c.get(f"{base}/quotes/{q['id']}").json()["status"] == "accepted"
    assert c.get(f"{base}/quotes").json()[0]["tours"] == 1

    # The calculation at 19 became the tour's budget.
    bud = c.get(f"{base}/tours/{r['tour_id']}/budget").json()
    assert bud["set"] and bud["tourists"] == 19 and bud["leaders"] == 1
    assert abs(bud["totals"]["planned"] - price["group_cost"]) < 0.5
    assert bud["totals"]["revenue_planned"] == price["group_revenue"] and bud["totals"]["revenue_actual"] == 0
    c.post(f"{base}/invoices/{r['invoice_id']}/issue")
    bud = c.get(f"{base}/tours/{r['tour_id']}/budget").json()
    assert bud["totals"]["revenue_actual"] == price["group_revenue"] + price["single_supplement"]

    assert c.post(f"{base}/quotes/{q['id']}/create-tour", json={"tourists": 40, "start_date": "2026-11-01"}).status_code == 422
    plain = mk(c, f"{base}/tours", {"code": "X", "start_date": "2026-11-01", "end_date": "2026-11-01"})
    none = c.get(f"{base}/tours/{plain['id']}/budget").json()
    assert none["has_quote"] is False and none["set"] is False and none["lines"] == []
    assert c.post(f"{base}/tours/{plain['id']}/budget", json={}).status_code == 409


def test_quotes_are_manager_only_and_scoped(make_client):
    owner, base, ws = agency(make_client)
    q = new_quote(owner, base)
    staff = make_client()
    signup(staff, "s@a.com")
    invite_and_join(owner, ws, staff, "s@a.com", "staff")
    assert staff.get(f"{base}/quotes").status_code == 403
    assert staff.get(f"{base}/quotes/{q['id']}").status_code == 403

    other, base_b, _ = agency(make_client, "b@b.com", "B")
    assert other.get(f"{base_b}/quotes/{q['id']}").status_code == 404
    assert other.post(f"{base_b}/quotes", json={"copy_of": q["id"]}).status_code == 404
    assert other.get(f"{base_b}/quotes/{q['id']}/price-list.xlsx").status_code == 404
    # B can't point its own quote at A's vendors or clients.
    mine = new_quote(other, base_b)
    hotel_a = mk(owner, f"{base}/vendors", {"name": "A hotel", "type": "hotel"})
    d = doc(mine)
    d["nights"][0]["vendor_id"] = hotel_a["id"]
    assert other.put(f"{base_b}/quotes/{mine['id']}", json=d).status_code == 422


def test_formula_check_and_explain_endpoints(make_client):
    c, base, _ = agency(make_client)
    d = doc(new_quote(c, base))
    r = c.post(f"{base}/quotes/preview/formula", json={"doc": d, "expression": "ceil(persons / 7)", "tourists": 19}).json()
    assert r["ok"] and [v["value"] for v in r["values"]] == [2, 3, 4]
    r = c.post(f"{base}/quotes/preview/formula", json={"doc": d, "expression": "dayz", "tourists": 19}).json()
    assert r["error_code"] == "unknown_name" and r["error_params"]["suggestion"] == "days"
    b = c.post(f"{base}/quotes/preview/breakdown", params={"tourists": 19}, json=d).json()
    assert b["steps"]["profit_rule"] == "group" and b["steps"]["rounding_mode"] == "floor"


def test_history_records_changes_and_restores(make_client):
    c, base, _ = agency(make_client)
    q = new_quote(c, base)
    assert q["version"] == 1
    url = f"{base}/quotes/{q['id']}"
    d = doc(q)
    bus = next(i for i in d["items"] if i["name"] == "Bus")
    bus["unit_price"] = "600"
    d["variables"]["armenia_meals"] = "2"
    d["nights"][5]["hotel_name"] = "Sheraton"
    assert c.put(url, params={"note": "new bus contract"}, json=d).json()["version"] == 2
    # Saving the same thing again doesn't add a version.
    assert c.put(url, json=d).json()["version"] == 2

    vs = c.get(f"{url}/versions").json()
    assert [v["number"] for v in vs] == [2, 1] and vs[0]["note"] == "new bus contract" and vs[0]["changes"] == 3
    assert vs[0]["created_by"] == "boss"

    diff = c.get(f"{url}/versions/2").json()
    got = {(x["area"], x["label"], x.get("field")): (x["before"], x["after"]) for x in diff["changes"]}
    assert got[("items", "Bus", "unit_price")] == ("550", "600")
    assert got[("variables", "armenia_meals", None)] == ("4", "2")
    assert got[("nights", "6. Batumi · Sheraton", "hotel_name")] == ("GreenWood 5*", "Sheraton")
    impact = {x["tourists"]: x for x in diff["impact"]}
    assert impact[19]["before"] == 1174 and impact[19]["after"] != 1174

    restored = c.post(f"{url}/versions/1/restore").json()
    assert restored["version"] == 3 and {t["tourists"]: t for t in restored["result"]["tiers"]}[19]["price_dbl"] == 1174
    assert c.get(f"{url}/versions").json()[0]["note"] == "restore:1"
    assert c.get(f"{url}/versions/9").status_code == 404


def test_scenarios_follow_the_calculation_and_compare(make_client):
    c, base, _ = agency(make_client)
    q = new_quote(c, base)
    url = f"{base}/quotes/{q['id']}"
    d = doc(q)
    wine = next(i for i in d["items"] if i["name"].startswith("Wine"))
    batumi = d["nights"][5]
    sc = mk(c, f"{url}/scenarios", {"name": "No Armenia, with wine", "patch": {
        "variables": {"armenia_meals": "0"}, "items": {wine["id"]: {"included": True}},
        "nights": {batumi["id"]: {"hotel_name": "Sheraton", "price_dbl": "140"}}}})
    kinds = {(x["area"], x["field"] or x["key"]) for x in sc["changes"]}
    assert ("variables", "armenia_meals") in kinds and ("items", "included") in kinds and ("nights", "price_dbl") in kinds
    assert sc["stale"] == []

    cmp = c.get(f"{url}/compare").json()
    base19 = {t["tourists"]: t for t in cmp["base"]["tiers"]}[19]["price_dbl"]
    sc19 = {t["tourists"]: t for t in cmp["scenarios"][0]["result"]["tiers"]}[19]["price_dbl"]
    assert base19 == 1174 and sc19 != base19

    # A change to the calculation itself flows into the scenario.
    d["profit_min_pp"] = "200"
    c.put(url, json=d)
    cmp2 = c.get(f"{url}/compare").json()
    assert {t["tourists"]: t for t in cmp2["scenarios"][0]["result"]["tiers"]}[25]["profit_pp"] == 200

    # Preview without saving; stale references are reported, not fatal.
    pv = c.post(f"{url}/scenarios/preview", json={"name": "x", "patch": {"settings": {"days": 10}}}).json()
    assert pv["changes"][0]["field"] == "days"
    d["items"] = [i for i in d["items"] if i["id"] != wine["id"]]
    c.put(url, json=d)
    assert c.get(f"{url}/scenarios").json()[0]["stale"] == [f"item:{wine['id']}"]

    # Separate calculation from a scenario, then apply it to the original.
    copy = mk(c, f"{base}/quotes", {"copy_of": q["id"], "scenario_id": sc["id"], "name": "HM · Sheraton"})
    assert copy["variables"]["armenia_meals"] == "0" and copy["nights"][5]["hotel_name"] == "Sheraton"
    applied = c.post(f"{url}/scenarios/{sc['id']}/apply").json()
    assert applied["variables"]["armenia_meals"] == "0"
    assert c.get(f"{url}/scenarios").json() == []
    assert c.get(f"{url}/versions").json()[0]["note"] == "scenario:No Armenia, with wine"


def test_history_and_scenarios_are_scoped(make_client):
    a, base_a, _ = agency(make_client, "a@a.com", "A")
    b, base_b, _ = agency(make_client, "b@b.com", "B")
    q = new_quote(a, base_a)
    sc = mk(a, f"{base_a}/quotes/{q['id']}/scenarios", {"name": "x"})
    assert b.get(f"{base_b}/quotes/{q['id']}/versions").status_code == 404
    assert b.get(f"{base_b}/quotes/{q['id']}/compare").status_code == 404
    assert b.post(f"{base_b}/quotes/{q['id']}/scenarios/{sc['id']}/apply").status_code == 404
    mine = new_quote(b, base_b)
    # A's scenario id can't be used through B's own quote either.
    assert b.delete(f"{base_b}/quotes/{mine['id']}/scenarios/{sc['id']}").status_code == 404
    hotel_a = mk(a, f"{base_a}/vendors", {"name": "A hotel", "type": "hotel"})
    night = doc(mine)["nights"][0]["id"]
    r = b.post(f"{base_b}/quotes/{mine['id']}/scenarios", json={"name": "y", "patch": {"nights": {night: {"vendor_id": hotel_a["id"]}}}})
    assert r.status_code == 422


def test_autosave_keeps_one_version_per_sitting_and_versions_can_be_named(make_client):
    c, base, _ = agency(make_client)
    q = new_quote(c, base)
    url = f"{base}/quotes/{q['id']}"
    d = doc(q)
    bus = next(i for i in d["items"] if i["name"] == "Bus")
    # The first autosave adds a version (the first one, how it began, is never changed)…
    bus["unit_price"] = "600"
    assert c.put(url, params={"autosave": True}, json=d).json()["version"] == 2
    # …the next ones of the same sitting update it.
    bus["unit_price"] = "610"
    assert c.put(url, params={"autosave": True}, json=d).json()["version"] == 2
    d["variables"]["armenia_meals"] = "2"
    assert c.put(url, params={"autosave": True}, json=d).json()["version"] == 2
    vs = c.get(f"{url}/versions").json()
    assert vs[0]["changes"] == 2
    assert c.get(f"{url}/versions/2").json()["changes"][0]["after"] in ("610", "2")
    # Ctrl+S (no autosave) is a version of its own; nothing new + a note names the latest.
    bus["unit_price"] = "620"
    assert c.put(url, json=d).json()["version"] == 3
    assert c.put(url, params={"note": "sent to client"}, json=d).json()["version"] == 3
    assert c.get(f"{url}/versions").json()[0]["note"] == "sent to client"
    # A named version is not changed by autosave.
    bus["unit_price"] = "630"
    assert c.put(url, params={"autosave": True}, json=d).json()["version"] == 4
    # Any version can be (re)named.
    assert c.patch(f"{url}/versions/2", json={"note": "first draft"}).status_code == 200
    assert next(v for v in c.get(f"{url}/versions").json() if v["number"] == 2)["note"] == "first draft"
    assert c.patch(f"{url}/versions/9", json={"note": "x"}).status_code == 404
    # Changed and changed back in one sitting: the version goes away again.
    bus["unit_price"] = "620"
    assert c.put(url, params={"autosave": True}, json=d).json()["version"] == 3


def test_a_free_service_is_free_not_missing(make_client):
    c, base, _ = agency(make_client)
    museum = mk(c, f"{base}/vendors", {"name": "Stalin Museum", "type": "activity"})
    paid = mk(c, f"{base}/vendors", {"name": "Wine Cellar", "type": "activity"})
    free_s = c.put(f"{base}/vendors/{museum['id']}/services", json=[
        {"name": "Entrance", "category": "ticket", "unit": "person", "rates": [{"price": "0", "currency": "GEL"}]}]).json()[0]
    paid_s = c.put(f"{base}/vendors/{paid['id']}/services", json=[
        {"name": "Tasting", "category": "activity", "unit": "person", "rates": [{"price": "45", "currency": "GEL"}]}]).json()[0]
    s = mk(c, f"{base}/series", {"code": "FR", "days": [
        {"city": "Gori", "extras": [{"service_id": free_s["id"], "quantity": None, "notes": ""},
                                    {"service_id": paid_s["id"], "quantity": None, "notes": ""}]}, {"city": "Tbilisi"}]})
    r = c.post(f"{base}/quotes/route-fill", json={"series_id": s["id"]}).json()
    entrance = next(i for i in r["items"] if "Entrance" in i["name"])
    tasting = next(i for i in r["items"] if "Tasting" in i["name"])
    assert entrance["free"] is True and float(entrance["unit_price"]) == 0
    assert tasting["free"] is False and float(tasting["unit_price"]) == 45
    assert not any("Entrance" in w for w in r["warnings"])
