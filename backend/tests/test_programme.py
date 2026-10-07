"""The programme to send: by day, in the chosen language, from what's on screen (saved or not)."""
import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import load_workbook

from app import programme
from app.config import settings

from .test_quotes_api import mk
from .test_tours import agency


class FakeTranslator:
    """Answers like Claude: the list, each item marked with the language. Counts the calls."""
    def __init__(self):
        self.calls = 0
        self.messages = self

    def create(self, model, max_tokens, messages):
        self.calls += 1
        prompt = messages[0]["content"]
        items = json.loads(prompt[prompt.index("["):])
        lang = "ZH" if "Chinese" in prompt else "KA" if "Georgian" in prompt else "EN"
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=json.dumps([f"{lang}:{x}" for x in items], ensure_ascii=False))])


def docx_text(content: bytes) -> str:
    """The Word file as text, in reading order: paragraphs, and a timeline table row as "time  what"."""
    def text(el):
        return "".join(t.text or "" for t in el.iter() if t.tag.endswith("}t"))
    out = []
    for block in Document(io.BytesIO(content)).element.body.iterchildren():
        if block.tag.endswith("}p"):
            out.append(text(block))
        elif block.tag.endswith("}tbl"):
            for tr in (x for x in block.iter() if x.tag.endswith("}tr")):
                cells = [c for c in tr if c.tag.endswith("}tc")]
                out.append("  ".join("\n".join(text(p) for p in c if p.tag.endswith("}p")) for c in cells).strip())
    return "\n".join(out)


@pytest.fixture(autouse=True)
def no_map(monkeypatch):
    """No real map in tests: a day without saved km stays without (one test gives its own map)."""
    monkeypatch.setattr(programme.geo, "legs", lambda db, cities, **kw: [])


@pytest.fixture
def translator(monkeypatch):
    monkeypatch.setattr(settings, "anthropic_api_key", "test-key")
    fake = FakeTranslator()
    monkeypatch.setattr(programme, "make_client", lambda ws: fake)
    return fake


def _days(hotel, resto, museum_service):
    return [
        {"city": "Tbilisi", "hotel_id": hotel, "dinner_id": resto, "schedule": [
            {"id": "a", "kind": "visit", "time": "15:00", "end": "17:00", "title": "Old Tbilisi walk"},
            {"id": "b", "kind": "meal", "time": "19:00", "ref": "dinner"}]},
        {"city": "Kazbegi", "stops": ["Mtskheta", "Ananuri"], "km": 160, "drive_minutes": [30, 40, 50], "drive_km": [25, 45, 90],
         "hotel_id": hotel, "lunch_mode": "own", "extras": [{"service_id": museum_service, "quantity": None, "notes": ""}],
         "notes": "Bring warm clothes"},
        {"city": "Tbilisi", "km": 155},
    ]


def _setup(c, base):
    hotel = mk(c, f"{base}/vendors", {"name": "Rooms Hotel", "type": "hotel"})
    resto = mk(c, f"{base}/vendors", {"name": "Barbarestan", "type": "restaurant"})
    museum = mk(c, f"{base}/vendors", {"name": "Gergeti Trinity", "type": "activity"})
    svc = c.put(f"{base}/vendors/{museum['id']}/services", json=[
        {"name": "Jeep up the hill", "category": "activity", "unit": "person", "rates": [{"price": "20", "currency": "GEL"}]}]).json()[0]
    return _days(hotel["id"], resto["id"], svc["id"])


def test_programme_by_day_in_chinese_translated_once(make_client, translator):
    c, base, _ = agency(make_client)
    days = _setup(c, base)
    body = {"title": "Georgia Highlights", "start_date": "2026-11-01", "days": days, "lang": "zh"}
    p = c.post(f"{base}/programme", json=body).json()
    assert p["labels"]["breakfast"] == "早餐" and p["warnings"] == []
    d1, d2, d3 = p["days"]
    assert d1["title"].startswith("第1天") and d1["hotel"] == "ZH:Rooms Hotel" and d1["date"] == "2026-11-01"
    assert d1["steps"][0] == {"time": "15:00", "end": "17:00", "text": "ZH:Old Tbilisi walk", "drive": None, "note": "", "kind": "visit"}
    assert d1["steps"][1]["text"] == "晚餐: ZH:Barbarestan"
    # Day 2: the road through its places, km and the coach time (×1.25: 120 min → 2 h 30 min).
    assert d2["route"] == ["ZH:Tbilisi", "ZH:Mtskheta", "ZH:Ananuri", "ZH:Kazbegi"]
    assert d2["km"] == 160 and d2["drive"] == "2 小时 30 分钟"
    assert ["早餐", "ZH:Rooms Hotel"] in d2["meals"] and ["午餐", "自理"] in d2["meals"]
    assert d2["activities"] == ["ZH:Gergeti Trinity: ZH:Jeep up the hill"] and d2["notes"] == "ZH:Bring warm clothes"
    assert d3["km"] == 155 and p["total_km"] == 315 and p["title"] == "ZH:Georgia Highlights"
    # No prices anywhere.
    assert "GEL" not in json.dumps(p["days"], ensure_ascii=False)
    calls = translator.calls
    c.post(f"{base}/programme", json=body)
    assert translator.calls == calls                     # kept: not asked (or paid) again


def test_georgian_texts_stay_and_files_download(make_client, translator):
    c, base, _ = agency(make_client)
    days = _setup(c, base)
    days[0]["city"] = "თბილისი"
    p = c.post(f"{base}/programme", json={"days": days, "lang": "ka", "include": {"schedule": False}}).json()
    assert p["days"][0]["route"] == ["თბილისი"] and p["days"][0]["steps"] == []
    assert p["days"][1]["route"][0] == "თბილისი"
    r = c.post(f"{base}/programme.docx", json={"title": "Test", "days": days, "lang": "en"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    text = docx_text(r.content)
    assert "Day 2" in text and "Overnight: Rooms Hotel" in text and "Distance: 160 km" in text
    x = c.post(f"{base}/programme.xlsx", json={"title": "Test", "days": days, "lang": "en"})
    ws = load_workbook(io.BytesIO(x.content)).active
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=6)]
    assert ws["A1"].value == "Test" and rows[0][0].startswith("Day 1")              # a coloured row for each day
    assert [None, "15:00–17:00", "Old Tbilisi walk", None] in rows                 # one row per step
    assert any(r[0] and r[0].startswith("Day 2") and r[3] == "Distance: 160 km · Driving: 2 h 30 min" for r in rows)


def test_without_an_ai_key_texts_stay_as_written(make_client, monkeypatch):
    monkeypatch.setattr(settings, "anthropic_api_key", "")
    c, base, _ = agency(make_client)
    p = c.post(f"{base}/programme", json={"days": _setup(c, base), "lang": "zh"}).json()
    assert p["warnings"] == ["no_ai_key"] and p["days"][0]["hotel"] == "Rooms Hotel" and p["labels"]["day"] == "第{n}天"


def test_place_to_place_km_and_only_some_days(make_client, translator):
    c, base, _ = agency(make_client)
    days = _setup(c, base)
    p = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False, "only_days": [2]}).json()
    assert [d["number"] for d in p["days"]] == [2] and p["total_km"] == 160
    legs = p["days"][0]["legs"]
    assert [(g["start"], g["end"], g["km"]) for g in legs] == [("Tbilisi", "Mtskheta", 25), ("Mtskheta", "Ananuri", 45), ("Ananuri", "Kazbegi", 90)]
    assert legs[0]["drive"] == "40 min"          # 30 min by car x 1.25, to 5 minutes
    assert p["days"][0]["meals"][0] == ["Breakfast", "Rooms Hotel"]     # day 1 left out, its hotel still counts
    text = docx_text(c.post(f"{base}/programme.docx", json={"days": days, "lang": "en"}).content)
    assert "Mtskheta \u2192 Ananuri \u00b7 45 km" in text


def test_distance_from_the_map_when_the_day_has_none_saved(make_client, monkeypatch):
    """The day plan shows the map's km live; the file must show the same, not only km saved on the day."""
    asked = []
    def fake_legs(db, cities, **kw):
        asked.append(list(cities))
        return [{"from": None, "to": cities[0], "km": 0, "minutes": 0}] + [
            {"from": a, "to": b, "km": 100, "minutes": 80} for a, b in zip(cities, cities[1:])]
    monkeypatch.setattr(programme.geo, "legs", fake_legs)
    c, base, _ = agency(make_client)
    days = [{"city": "Tbilisi"}, {"city": "Kutaisi", "stops": ["Gori"]}]
    p = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False}).json()
    d2 = p["days"][1]
    assert asked == [["Tbilisi", "Gori", "Kutaisi"]]
    assert d2["km"] == 200 and d2["drive"] == "3 h 20 min" and len(d2["legs"]) == 2 and p["total_km"] == 200
    assert {k: d2["legs"][0][k] for k in ("start", "end", "km", "drive")} == {"start": "Tbilisi", "end": "Gori", "km": 100, "drive": "1 h 40 min"}


def test_distance_between_the_destinations_of_the_day_programme(make_client, monkeypatch):
    """With a timed programme the file shows each drive between its destinations, like the day plan."""
    asked = []
    def fake_legs(db, cities, **kw):
        asked.append(list(cities))
        return [{"from": None, "to": cities[0], "km": 0, "minutes": 0}] + [
            {"from": a, "to": b, "km": 10 * (j + 1), "minutes": 12} for j, (a, b) in enumerate(zip(cities, cities[1:]))]
    monkeypatch.setattr(programme.geo, "legs", fake_legs)
    c, base, _ = agency(make_client)
    lodge = mk(c, f"{base}/vendors", {"name": "Gudauri Lodge", "type": "hotel", "city": "Gudauri", "address": "Gudauri ski resort"})
    jvari = mk(c, f"{base}/vendors", {"name": "Jvari Monastery", "type": "activity", "city": "Mtskheta", "address": "Jvari, Mtskheta"})
    days = [{"city": "Tbilisi"}, {"city": "Gudauri", "hotel_id": lodge["id"], "schedule": [
        {"id": "1", "kind": "visit", "time": "10:00", "title": "Jvari Monastery"},
        {"id": "2", "kind": "visit", "time": "12:00", "title": "Ananuri"},
        {"id": "3", "kind": "checkin", "time": "16:00", "ref": "hotel"}]}, {"city": "Tbilisi"}]
    p = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False}).json()
    d2 = p["days"][1]
    assert ["Tbilisi", "Jvari, Mtskheta", "Ananuri", "Gudauri ski resort"] in asked      # where the map looks
    # In the timeline: each drive just before the place it leads to.
    assert [(s["text"], s["drive"]["start"], s["drive"]["km"]) for s in d2["steps"]] == [
        ("Jvari Monastery", "Tbilisi", 10), ("Ananuri", "Jvari Monastery", 20), ("Check-in: Gudauri Lodge", "Ananuri", 30)]
    assert d2["km"] == 60 and d2["title"] == "Day 2 \u00b7 Tbilisi \u2192 Gudauri"


def test_one_timeline_drive_before_each_place(make_client, monkeypatch):
    monkeypatch.setattr(programme.geo, "legs", lambda db, cities, **kw: [{"from": None, "to": cities[0], "km": 0, "minutes": 0}] + [
        {"from": a, "to": b, "km": 10 * (j + 1), "minutes": 32} for j, (a, b) in enumerate(zip(cities, cities[1:]))])
    c, base, _ = agency(make_client)
    days = [{"city": "Tbilisi"}, {"city": "Gudauri", "schedule": [
        {"id": "1", "kind": "visit", "time": "10:00", "end": "11:00", "title": "Jvari Monastery"},
        {"id": "2", "kind": "visit", "time": "12:00", "title": "Ananuri"}]}, {"city": "Tbilisi"}]
    p = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False}).json()
    d2 = p["days"][1]
    assert d2["legs"] == []                                   # no separate list: it's in the timeline
    assert [(s["text"], s["drive"] and s["drive"]["km"]) for s in d2["steps"]] == [
        ("Jvari Monastery", 10), ("Ananuri", 20), ("Gudauri", 30)]
    assert d2["steps"][0]["drive"]["drive"] == "40 min" and d2["steps"][2]["time"] == "12:40"     # arrives after the drive
    # Without the timed programme the place-to-place list stays.
    q = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False, "include": {"schedule": False}}).json()
    assert q["days"][1]["steps"] == [] and len(q["days"][1]["legs"]) == 3
    text = docx_text(c.post(f"{base}/programme.docx", json={"days": days, "lang": "en", "translate": False}).content)
    assert "09:20\u201310:00  \u2193 10 km \u00b7 40 min\n10:00\u201311:00  Jvari Monastery" in text
    assert "11:00\u201312:00  \u2193 20 km \u00b7 40 min\n12:00  Ananuri" in text


def test_flights_in_the_timeline_with_the_airports_name(make_client, monkeypatch):
    asked = []
    def fake_legs(db, cities, **kw):
        asked.append(list(cities))
        return [{"from": None, "to": cities[0], "km": 0, "minutes": 0}] + [
            {"from": a, "to": b, "km": 18, "minutes": 24} for a, b in zip(cities, cities[1:])]
    monkeypatch.setattr(programme.geo, "legs", fake_legs)
    c, base, _ = agency(make_client)
    hotel = mk(c, f"{base}/vendors", {"name": "Rooms Hotel", "type": "hotel", "city": "Tbilisi", "address": "Kostava 14, Tbilisi"})
    days = [
        {"city": "Tbilisi", "hotel_id": hotel["id"], "flights": [{"direction": "arrival", "number": "TK 384", "time": "10:30", "airport": "tbs", "other": "Istanbul"}],
         "schedule": [{"id": "f", "kind": "flight", "time": "10:30", "ref": "flight:0"}, {"id": "h", "kind": "checkin", "time": "14:00", "ref": "hotel"}]},
        {"city": "Tbilisi", "flights": [{"direction": "departure", "number": "TK 385", "time": "17:00", "airport": "TBS"}],
         "schedule": [{"id": "g", "kind": "flight", "time": "17:00", "ref": "flight:0"}]},
    ]
    p = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False}).json()
    d1, d2 = p["days"]
    assert d1["steps"][0]["text"] == "Arrival TK384 (Istanbul) \u00b7 Tbilisi International Airport" and d1["steps"][0]["time"] == "10:30"
    assert d1["steps"][1]["drive"]["start"] == "Tbilisi International Airport" and d1["steps"][1]["drive"]["km"] == 18
    assert d1["flights"] == []                                  # in the timeline, not twice
    assert ["Tbilisi International Airport", "Kostava 14, Tbilisi"] in asked
    assert d2["steps"][-1]["text"].startswith("Departure TK385") and d2["steps"][-1]["drive"]["end"] == "Tbilisi International Airport"
    z = c.post(f"{base}/programme", json={"days": days, "lang": "zh", "translate": False, "include": {"schedule": False}}).json()
    assert z["days"][0]["flights"] == ["10:30 \u62b5\u8fbe TK384 (Istanbul) \u00b7 \u7b2c\u6bd4\u5229\u65af\u56fd\u9645\u673a\u573a"]


def test_no_drive_after_dinner_at_the_nights_hotel(make_client, monkeypatch):
    asked = []
    def fake_legs(db, cities, **kw):
        asked.append(list(cities))
        return [{"from": None, "to": cities[0], "km": 0, "minutes": 0}] + [
            {"from": a, "to": b, "km": 8, "minutes": 16} for a, b in zip(cities, cities[1:])]
    monkeypatch.setattr(programme.geo, "legs", fake_legs)
    c, base, _ = agency(make_client)
    hotel = mk(c, f"{base}/vendors", {"name": "Marco Polo Hotel Gudauri", "type": "hotel", "city": "Gudauri",
                                      "address": "Marco Polo Hotel Gudauri, Gudauri"})
    days = [{"city": "Tbilisi"}, {"city": "Gudauri", "hotel_id": hotel["id"], "dinner_id": hotel["id"], "schedule": [
        {"id": "v", "kind": "visit", "time": "18:35", "end": "18:50", "title": "Gudauri View Point"},
        {"id": "d", "kind": "meal", "time": "19:30", "ref": "dinner"}]}, {"city": "Tbilisi"}]
    p = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False}).json()
    steps = p["days"][1]["steps"]
    # View point → (drive) → dinner at the hotel, and the day ends there: no drive "to Gudauri".
    assert [s["text"] for s in steps] == ["Gudauri View Point", "Dinner: Marco Polo Hotel Gudauri"]
    assert ["Tbilisi", "Gudauri View Point", "Marco Polo Hotel Gudauri, Gudauri"] in asked


def test_the_menu_under_the_restaurant(make_client):
    c, base, _ = agency(make_client)
    resto = mk(c, f"{base}/vendors", {"name": "Barbarestan", "type": "restaurant", "city": "Tbilisi"})
    menus = c.put(f"{base}/vendors/{resto['id']}/services", json=[
        {"code": "dinner", "name": "Dinner", "category": "meal", "unit": "person", "dishes": ["Khinkali", "Mtsvadi"],
         "rates": [{"price": "30", "currency": "GEL"}]},
        {"name": "Barbarestan · Lunch 1", "category": "meal", "unit": "person", "dishes": ["Lobio", "Mchadi"],
         "rates": [{"price": "25", "currency": "GEL"}]}]).json()
    lunch_menu = next(m for m in menus if m["name"].endswith("Lunch 1"))
    days = [{"city": "Tbilisi", "lunch_id": resto["id"], "lunch_menu_id": lunch_menu["id"], "dinner_id": resto["id"], "schedule": [
        {"id": "l", "kind": "meal", "time": "13:00", "ref": "lunch"}, {"id": "d", "kind": "meal", "time": "19:30", "ref": "dinner"}]}]
    p = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False}).json()
    d = p["days"][0]
    assert [(s["text"], s["note"]) for s in d["steps"]] == [("Lunch: Barbarestan", "Lobio, Mchadi"), ("Dinner: Barbarestan", "Khinkali, Mtsvadi")]
    assert ["Lunch", "Barbarestan"] in d["meals"]            # the timeline has the menu: not twice
    q = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False, "include": {"schedule": False}}).json()
    assert ["Lunch", "Barbarestan: Lobio, Mchadi"] in q["days"][0]["meals"] and ["Dinner", "Barbarestan: Khinkali, Mtsvadi"] in q["days"][0]["meals"]
    text = docx_text(c.post(f"{base}/programme.docx", json={"days": days, "lang": "en", "translate": False}).content)
    assert "13:00  Lunch: Barbarestan\nLobio, Mchadi" in text


def test_each_drive_has_its_own_time(make_client, monkeypatch):
    monkeypatch.setattr(programme.geo, "legs", lambda db, cities, **kw: [{"from": None, "to": cities[0], "km": 0, "minutes": 0}] + [
        {"from": a, "to": b, "km": 86, "minutes": 72} for a, b in zip(cities, cities[1:])])
    c, base, _ = agency(make_client)
    days = [{"city": "Tbilisi"}, {"city": "Kutaisi", "schedule": [
        {"id": "b", "kind": "breakfast", "time": "08:00", "end": "08:30", "ref": "breakfast"},
        {"id": "g", "kind": "visit", "time": "11:00", "end": "12:30", "title": "Gori"}]}, {"city": "Tbilisi"}]
    p = c.post(f"{base}/programme", json={"days": days, "lang": "en", "translate": False}).json()
    gori, arrive = p["days"][1]["steps"][-2:]
    # From when breakfast ends to when Gori starts (the drive itself is 72 car minutes x 1.25 = 90).
    assert (gori["drive"]["depart"], gori["drive"]["arrive"], gori["drive"]["minutes"]) == ("08:30", "11:00", 90)
    assert gori["kind"] == "visit" and arrive["kind"] == "arrive" and arrive["text"] == "Kutaisi"
    assert (arrive["drive"]["depart"], arrive["drive"]["arrive"], arrive["time"]) == ("12:30", "14:00", "14:00")
