"""The agency's own kinds of cost (Flight, Train, Visa…) and ticket lines with their journey: legs,
taxes on top of each fare, the date tickets must be issued by."""
from datetime import date
from decimal import Decimal as D

import pytest

from app.quote_engine import breakdown, compute
from app.quote_schemas import CalcStyle, CostKind, PersonType, PriceStep, QuoteIn, QuoteItemIn, QuoteNightIn, TravelInfo, TravelLeg, Usage

ADULTS, KIDS, LEADER = "adults", "kids", "leader"


def flight(**kw) -> QuoteItemIn:
    base = dict(name="Flight TBS-EVN-TBS", category="transport", kind="flight", currency="USD", unit_price=D(150),
                basis="per_person", who="persons",
                travel=TravelInfo(mode="flight", taxes=D(40), ticket_by=date(2026, 9, 1),
                                  legs=[TravelLeg(from_place="Tbilisi", to_place="Yerevan", on=date(2026, 10, 12), number="A9 123"),
                                        TravelLeg(from_place="Yerevan", to_place="Tbilisi", on=date(2026, 10, 15))]))
    base.update(kw)
    return QuoteItemIn(**base)


def trip(items, recipe=()) -> QuoteIn:
    return QuoteIn(
        name="Fly", currency="USD", days=4, pax_min=10, pax_max=20,
        people=[PersonType(id=ADULTS, name="Adults"), PersonType(id=KIDS, name="Children", count=2, room="none"),
                PersonType(id=LEADER, name="Leader", role="free", count=1, room="sgl")],
        nights=[QuoteNightIn(currency="USD", price_dbl=D(100), price_sgl=D(80))],
        items=items, recipe=list(recipe))


def tier(q, n=10):
    return next(t for t in compute(q).tiers if t.tourists == n)


def test_taxes_go_on_every_fare_including_a_child_fare():
    q = trip([flight(uses={KIDS: Usage(mode="own", price=D(100))})])
    b = breakdown(q, 10)
    item = b.items[0]
    # 10 adults × (150 + 40) + 2 children × (100 + 40) + the leader (150 + 40)
    assert item.amount == pytest.approx(10 * 190 + 2 * 140 + 190)
    assert item.unit_price == pytest.approx(190) and item.kind == "flight"


def test_a_recipe_can_mark_up_the_agencys_own_kind_alone():
    plain = tier(trip([flight()]))
    marked = tier(trip([flight()], [PriceStep(kind="markup", value=D(10), cats=["flight"])]))
    adults = plain.types[0]
    # Only the flights (own fare + share of the leader's ticket) get the 10%, not the hotel.
    flight_pp = 190 + 190 / 12
    assert marked.types[0].price == pytest.approx(adults.cost + flight_pp * 0.10, abs=0.02)
    # A markup on "transport" (the base) leaves the agency's own kind alone.
    other = tier(trip([flight()], [PriceStep(kind="markup", value=D(10), cats=["transport"])]))
    assert other.types[0].price == pytest.approx(adults.cost, abs=0.01)


def test_kinds_are_the_agencys_and_survive_saving_our_way(make_client):
    from .test_tours import agency
    c, base, _ = agency(make_client)
    kinds = [{"id": "flight", "name": "Flight", "base": "transport", "icon": "plane", "color": "sky", "travel": True},
             {"id": "visa", "name": "Visa", "base": "other", "icon": "file", "color": "violet"}]
    assert c.put(f"{base}/calc-style/kinds", json=kinds).status_code == 200
    # "Our way" saved without kinds keeps them.
    style = {"people": [{"id": "a", "name": "Adults", "role": "pay", "count": None, "room": "dbl"}], "recipe": []}
    assert c.put(f"{base}/calc-style", json=style).status_code == 200
    got = c.get(f"{base}/calc-style").json()
    assert [k["id"] for k in got["kinds"]] == ["flight", "visa"] and got["kinds"][0]["travel"] is True
    # A built-in name or a duplicate is refused.
    assert c.put(f"{base}/calc-style/kinds", json=[{"id": "hotel", "name": "Hotel"}]).status_code == 422
    assert c.put(f"{base}/calc-style/kinds", json=[kinds[0], kinds[0]]).status_code == 422
    # A calculation with a ticket line saves and prices.
    q = c.post(f"{base}/quotes", json={"preset": "blank"}).json()
    doc = {k: v for k, v in q.items() if k not in ("id", "updated_at", "version", "result")}
    doc["items"] = [flight().model_dump(mode="json")]
    doc["fx"] = {**doc.get("fx", {}), "USD": "1"} if doc["currency"] != "USD" else doc.get("fx", {})
    r = c.put(f"{base}/quotes/{q['id']}", json=doc)
    assert r.status_code == 200, r.text
    assert r.json()["items"][0]["travel"]["legs"][0]["to_place"] == "Yerevan"


def test_style_without_kinds_still_validates():
    assert CalcStyle.model_validate({}).kinds == []
    with pytest.raises(ValueError):
        CostKind(id="Bad Id", name="x")


def test_a_tour_from_the_calculation_gets_its_flights_and_a_ticket_task(make_client):
    from .test_quotes_api import doc as editable, new_quote
    from .test_tours import agency, mk
    c, base, _ = agency(make_client)
    q = new_quote(c, base)
    d = editable(q)
    fl = flight(day=1, travel=TravelInfo(mode="flight", taxes=D(40), ticket_by=date(2026, 9, 20), baggage="23 kg",
                                         legs=[TravelLeg(from_place="Istanbul", to_place="Tbilisi", number="tk 378", arrive="06:10"),
                                               TravelLeg(from_place="Tbilisi", to_place="Istanbul", number="TK379", depart="07:35")]))
    d["items"].append(fl.model_dump(mode="json"))
    assert c.put(f"{base}/quotes/{q['id']}", json=d).status_code == 200
    r = mk(c, f"{base}/quotes/{q['id']}/create-tour", {"tourists": 19, "start_date": "2026-11-01"})
    t = c.get(f"{base}/tours/{r['tour_id']}").json()
    first, last = t["days"][0]["flights"], t["days"][-1]["flights"]
    assert first == [{"direction": "arrival", "number": "TK378", "time": "06:10", "airport": "", "other": "Istanbul", "notes": ""}]
    assert last[0]["direction"] == "departure" and last[0]["number"] == "TK379" and last[0]["time"] == "07:35"
    tasks = [x for x in c.get(f"{base}/tasks?who=all").json() if x["title"].startswith("Issue tickets")]
    assert len(tasks) == 1 and tasks[0]["due_date"] == "2026-09-20" and "Istanbul → Tbilisi" in tasks[0]["title"]
    lines = c.get(f"{base}/tours/{r['tour_id']}/budget").json()["lines"]
    ticket = next(x for x in lines if x["name"].startswith("Flight TBS-EVN-TBS"))
    assert "Istanbul → Tbilisi · Tbilisi → Istanbul" in ticket["name"] and "issue tickets by 2026-09-20" in ticket["notes"]
    # Giving the calculation to the tour again adds nothing twice.
    c.post(f"{base}/quotes/{q['id']}/apply", json={"tour_ids": [r["tour_id"]], "tourists": 19})
    assert len(c.get(f"{base}/tours/{r['tour_id']}").json()["days"][0]["flights"]) == 1
    assert len([x for x in c.get(f"{base}/tasks?who=all").json() if x["title"].startswith("Issue tickets")]) == 1


def test_a_line_can_be_a_share_of_what_everything_costs():
    bus = QuoteItemIn(name="Bus", category="transport", currency="USD", unit_price=D(1000))
    plain = breakdown(trip([flight(), bus]), 10)
    insurance = QuoteItemIn(name="Insurance 3%", category="other", currency="USD", unit_price=D(0), price_formula="cost_total * 0.03")
    b = breakdown(trip([flight(), bus, insurance]), 10)
    costs = sum(x.amount for x in plain.items) + plain.steps.hotel_dbl_pp * 10 + plain.steps.leader_room
    assert b.items[2].amount == pytest.approx(costs * 0.03, abs=0.01)
    # By kind too: 10% of the flights only.
    fee = QuoteItemIn(name="Ticketing fee", category="other", currency="USD", unit_price=D(0), price_formula="cost_flight * 0.10")
    assert breakdown(trip([flight(), bus, fee]), 10).items[2].amount == pytest.approx(plain.items[0].amount * 0.10, abs=0.01)


def test_a_line_that_comes_with_another_follows_it():
    dinner = QuoteItemIn(id="ktw", name="KTW dinner", category="meal", currency="USD", unit_price=D(30), basis="per_person")
    show = QuoteItemIn(name="Dance show", category="activity", currency="USD", unit_price=D(10), basis="per_person", with_line="ktw")
    on = breakdown(trip([dinner, show]), 10)
    assert on.items[1].amount == pytest.approx(10 * 13)      # 10 adults + 2 children + the leader
    off = breakdown(trip([dinner.model_copy(update={"included": False}), show]), 10)
    assert off.items[1].amount == 0 and off.tier.cost_dbl < on.tier.cost_dbl


def test_the_formula_library_is_the_agencys(make_client):
    from .test_tours import agency
    c, base, _ = agency(make_client)
    lib = [{"name": "Insurance 3%", "formula": "cost_total * 0.03", "note": "of all costs"}]
    assert c.put(f"{base}/calc-style/formulas", json=lib).status_code == 200
    c.put(f"{base}/calc-style", json={"people": [], "recipe": []})
    assert c.get(f"{base}/calc-style").json()["formulas"][0]["formula"] == "cost_total * 0.03"


def test_the_whole_calculation_downloads_as_a_workbook(make_client):
    import io
    from openpyxl import load_workbook
    from .test_quotes_api import new_quote
    from .test_tours import agency
    c, base, _ = agency(make_client)
    q = new_quote(c, base)
    for lang, first in (("en", "Prices"), ("ka", "ფასები")):
        r = c.get(f"{base}/quotes/{q['id']}/calculation.xlsx?lang={lang}&tourists=19")
        assert r.status_code == 200 and "calculation" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert [s.title for s in wb.worksheets][0] == first and "19+1" in wb.sheetnames and len(wb.worksheets) == 5
    prices = load_workbook(io.BytesIO(c.get(f"{base}/quotes/{q['id']}/calculation.xlsx").content)).worksheets[0]
    by_label = {row[0]: row for row in prices.iter_rows(values_only=True) if row and isinstance(row[0], str) and "+" in row[0]}
    t19 = {t["tourists"]: t for t in q["result"]["tiers"]}[19]
    assert by_label["19+1"][4] == t19["price_dbl"]
