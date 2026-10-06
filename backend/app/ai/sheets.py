"""Spreadsheets for the assistant: any layout, read as rows of text.

The assistant gets a short preview of each sheet in the message and pages through the rest with
the read_sheet tool, so a 3000-row price list doesn't fill its context at once.
"""
import csv
import io
from datetime import date, datetime

from openpyxl import load_workbook

MAX_ROWS = 5000          # per sheet
MAX_COLS = 40
PREVIEW_ROWS = 40


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return " ".join(str(v).split())[:300]


def _trim(rows: list[list[str]]) -> list[list[str]]:
    rows = [r[:MAX_COLS] for r in rows]
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    return [r[:width] for r in rows]


def read(filename: str, data: bytes) -> dict[str, list[list[str]]]:
    """{sheet name: rows}. Raises ValueError for anything that isn't a readable sheet."""
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".tsv"):
        text = None
        for enc in ("utf-8-sig", "cp1251", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel_tab if name.endswith(".tsv") else csv.excel
        rows = [[_cell(c) for c in r] for _, r in zip(range(MAX_ROWS), csv.reader(io.StringIO(text), dialect))]
        return {"Sheet1": _trim(rows)}
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises many kinds for broken files
        raise ValueError("Can't read this spreadsheet — save it as .xlsx or .csv") from e
    out = {}
    for ws in wb.worksheets:
        rows = [[_cell(c) for c in r] for _, r in zip(range(MAX_ROWS), ws.iter_rows(values_only=True))]
        rows = _trim(rows)
        if rows:
            out[ws.title] = rows
    wb.close()
    if not out:
        raise ValueError("The spreadsheet is empty")
    return out


def as_text(rows: list[list[str]], start: int = 1) -> str:
    """Rows as tab-separated lines, each starting with its row number (1-based)."""
    return "\n".join(f"{start + i}\t" + "\t".join(r) for i, r in enumerate(rows))


def preview(filename: str, sheets: dict[str, list[list[str]]]) -> str:
    parts = [f"Spreadsheet {filename!r}:"]
    for title, rows in sheets.items():
        more = f" (showing rows 1–{PREVIEW_ROWS}; read the rest with read_sheet)" if len(rows) > PREVIEW_ROWS else ""
        parts.append(f"\nSheet {title!r}: {len(rows)} rows{more}\n" + as_text(rows[:PREVIEW_ROWS]))
    return "\n".join(parts)
