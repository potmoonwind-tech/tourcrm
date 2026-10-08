"""Spreadsheet import for participants, tours, vendors and clients.

Partner agencies send lists in whatever shape they have, so parsing is forgiving:
- .xlsx or .csv (any common delimiter, UTF-8 or Windows-1251);
- the header row is found automatically, even under a title line;
- column names are matched in English, Georgian, Russian and Chinese;
- dates as Excel dates, 2026-09-06, 06.09.2026, 06/09/2026 or 06SEP2026.

Every import runs twice: once as a preview (commit=False) that reports what
would happen row by row, then for real. Invalid rows are skipped, never guessed.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from .deps import WorkspaceCtx
from .models import Client, Participant, Series, Tour, Vendor
from .routers.people import next_position
from .routers.tours import MAX_TOUR_DAYS, _code_taken, _generate_code, sync_days
from .schemas import ImportIssue, ImportResult

MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024


class ImportFileError(ValueError):
    """The file as a whole can't be read (wrong format, empty, too big)."""


# --- reading ---------------------------------------------------------------------------

class Table(list):
    """The rows of the first sheet, plus what a list from a partner often carries besides values:
    `merged` - for each cell inside a merged block, the block's (top row, rows it spans) (two people
    in one "twin" cell share a room); `others` - the other sheets' rows (passports on Sheet2)."""
    def __init__(self, rows=(), merged=None, others=None):
        super().__init__(rows)
        self.merged: dict[tuple[int, int], tuple[int, int]] = merged or {}
        self.others: list[list[tuple[int, list]]] = others or []


def _sheet_rows(ws, merged: dict | None) -> list[tuple[int, list]]:
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        rows.append((i, list(r)))
        if i > MAX_ROWS + 50:      # a tiny file can unpack to millions of rows: stop early
            break
    if merged is not None:
        for rng in ws.merged_cells.ranges:
            top, span = rng.min_row, rng.max_row - rng.min_row + 1
            for r in range(rng.min_row, min(rng.max_row, len(rows)) + 1):
                for c in range(rng.min_col - 1, rng.max_col):
                    merged[(r, c)] = (top, span)
    return rows


def read_table(data: bytes, filename: str) -> list[tuple[int, list]]:
    """[(spreadsheet row number, [cell values])], empty rows dropped (a Table: merged cells and other sheets too)."""
    if len(data) > MAX_BYTES:
        raise ImportFileError("File is larger than 5 MB")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(data), data_only=True)
        except Exception as e:  # corrupt / not really xlsx
            raise ImportFileError("Could not read the Excel file") from e
        merged: dict = {}
        rows = _sheet_rows(wb.worksheets[0], merged)
        others = [[(i, r) for i, r in _sheet_rows(ws, None) if any(_blank(v) is not None for v in r)] for ws in wb.worksheets[1:4]]
        wb.close()
    elif name.endswith((".csv", ".txt")):
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportFileError("Could not decode the CSV file")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [(i, r) for i, r in enumerate(csv.reader(io.StringIO(text), dialect), 1)]
    else:
        raise ImportFileError("Upload an .xlsx or .csv file")
    rows = [(i, r) for i, r in rows if any(_blank(v) is not None for v in r)]
    if not rows:
        raise ImportFileError("The file is empty")
    if name.endswith((".xlsx", ".xlsm")):
        return Table(rows, merged, [o for o in others if o])
    return Table(rows)


def _blank(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def norm_header(h) -> str:
    return re.sub(r"[\s_\-.*:#()/№'\"]+", "", str(h or "").lower())


# --- fields & value parsing ------------------------------------------------------------

@dataclass
class F:
    key: str
    aliases: tuple[str, ...]
    parse: Callable[[object], object] = lambda v: str(v).strip()
    required: bool = False
    example: object = ""


def as_text(limit: int) -> Callable[[object], str]:
    def p(v):
        if isinstance(v, float) and v.is_integer():
            v = int(v)  # phone / passport numbers stored as numbers in Excel
        s = str(v).strip()
        if len(s) > limit:
            raise ValueError(f"longer than {limit} characters")
        return s
    return p


_DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%Y.%m.%d",
                 "%d %b %Y", "%d %B %Y", "%d%b%Y", "%d.%m.%y", "%d/%m/%y")


def as_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 1000 < v < 80000:  # Excel serial day in a CSV
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip().split(" 00:00")[0]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{s}' is not a date (use 2026-09-06 or 06.09.2026)")


def as_int(v) -> int:
    try:
        n = int(float(str(v).strip()))
    except ValueError:
        raise ValueError(f"'{v}' is not a number") from None
    if n < 0:
        raise ValueError("must not be negative")
    return n


def as_choice(mapping: dict[str, str], label: str) -> Callable[[object], str]:
    table = {norm_header(k): val for k, val in mapping.items()}

    def p(v):
        key = norm_header(v)
        if key in table:
            return table[key]
        raise ValueError(f"unknown {label} '{v}'")
    return p


GENDER = as_choice({
    "m": "M", "male": "M", "man": "M", "mr": "M", "მ": "M", "მამრ": "M", "კაცი": "M", "м": "M", "муж": "M", "мужской": "M", "男": "M",
    "f": "F", "female": "F", "woman": "F", "mrs": "F", "ms": "F", "ქ": "F", "მდედრ": "F", "ქალი": "F", "ж": "F", "жен": "F", "женский": "F", "女": "F",
}, "gender")

_ROOM = as_choice({
    "sgl": "SGL", "single": "SGL", "sngl": "SGL", "1": "SGL", "单人间": "SGL", "单间": "SGL", "ერთადგილიანი": "SGL",
    "dbl": "DBL", "double": "DBL", "dbl room": "DBL", "2": "DBL", "双人间": "DBL", "大床": "DBL", "大床房": "DBL", "ორადგილიანი": "DBL",
    "twn": "TWN", "twin": "TWN", "标间": "TWN", "双床": "TWN", "双床房": "TWN",
    "trp": "TRP", "tpl": "TRP", "triple": "TRP", "3": "TRP", "三人间": "TRP", "სამადგილიანი": "TRP",
    "quad": "QUAD", "qdr": "QUAD", "4": "QUAD", "四人间": "QUAD",
    "dbl/twn": "DBL", "双人": "DBL", "标准间": "TWN", "单人": "SGL", "三人": "TRP",
    # Other languages partners write in: Russian, Spanish, Italian, German, French, Turkish.
    "двухместный": "DBL", "дабл": "DBL", "твин": "TWN", "одноместный": "SGL", "сингл": "SGL", "трехместный": "TRP",
    "doble": "DBL", "matrimonial": "DBL", "individual": "SGL", "sencilla": "SGL", "triple": "TRP",
    "doppia": "DBL", "matrimoniale": "DBL", "singola": "SGL", "tripla": "TRP",
    "doppelzimmer": "DBL", "einzelzimmer": "SGL", "zweibett": "TWN", "chambre double": "DBL", "chambre simple": "SGL",
    "çift": "DBL", "tek": "SGL",
}, "room type")


def ROOM(v) -> str:
    """ "TWN", or "标间/twin", "Twin (2 beds)": the first part that names a room type."""
    try:
        return _ROOM(v)
    except ValueError:
        for part in re.split(r"[/|,;()（）\s]+", str(v)):
            if part.strip():
                try:
                    return _ROOM(part)
                except ValueError:
                    continue
        raise


LEADER_WORDS = {norm_header(x) for x in ("领队", "tl", "tour leader", "leader", "escort", "ტურლიდერი", "лидер", "руководитель", "руководитель группы")}

STATUS = as_choice({
    "confirmed": "confirmed", "ok": "confirmed", "yes": "confirmed", "დადასტურებული": "confirmed", "подтвержден": "confirmed",
    "inquiry": "inquiry", "option": "inquiry", "request": "inquiry", "tentative": "inquiry", "მოთხოვნა": "inquiry", "запрос": "inquiry",
    "cancelled": "cancelled", "canceled": "cancelled", "cancel": "cancelled", "cxl": "cancelled", "გაუქმებული": "cancelled", "отмена": "cancelled",
}, "status")

VENDOR_TYPE = as_choice({
    "hotel": "hotel", "სასტუმრო": "hotel", "гостиница": "hotel", "отель": "hotel", "酒店": "hotel",
    "restaurant": "restaurant", "რესტორანი": "restaurant", "ресторан": "restaurant", "餐厅": "restaurant",
    "transport": "transport", "bus": "transport", "ტრანსპორტი": "transport", "транспорт": "transport", "车": "transport",
    "guide": "guide", "გიდი": "guide", "гид": "guide", "导游": "guide",
    "driver": "driver", "მძღოლი": "driver", "водитель": "driver", "司机": "driver",
    "activity": "activity", "აქტივობა": "activity", "other": "other", "სხვა": "other",
}, "vendor type")

TIMING = as_choice({"before": "before", "prepay": "before", "prepaid": "before", "წინასწარ": "before", "до": "before",
                    "after": "after", "postpay": "after", "შემდეგ": "after", "после": "after"}, "payment timing")

CLIENT_TYPE = as_choice({"agency": "agency", "სააგენტო": "agency", "агентство": "agency", "旅行社": "agency",
                         "direct": "direct", "პირდაპირი": "direct", "частный": "direct"}, "client type")


def as_currency(v) -> str:
    s = str(v).strip().upper()
    if len(s) != 3 or not s.isalpha():
        raise ValueError(f"'{v}' is not a 3-letter currency code")
    return s


# Shared aliases
NAME = ("name", "სახელი", "название", "имя", "名称")
NOTES = ("notes", "note", "remarks", "remark", "comment", "comments", "შენიშვნა", "შენიშვნები", "примечание", "комментарий", "备注")
PHONE = ("phone", "mobile", "tel", "telephone", "contact", "ტელეფონი", "მობილური", "телефон", "电话", "手机", "联系方式", "联系电话")
EMAIL = ("email", "e-mail", "mail", "ელფოსტა", "почта", "邮箱")

SPECS: dict[str, list[F]] = {
    "participants": [
        F("full_name", ("fullname", "full name", "name", "passenger", "guest", "tourist", "traveller", "სახელიგვარი",
                        "სახელი და გვარი", "სახელი", "ტურისტი", "фио", "имя", "姓名", "名字"), as_text(200), example="WANG Li"),
        # The passport name in Latin letters - one column, or surname and given name side by side under one header.
        F("en_name", ("english name", "name in english", "latin name", "passport name", "英文名", "英文姓名", "拼音", "拼音名",
                      "имя латиницей", "ლათინურად"), as_text(200)),
        F("local_name", ("中文名", "中文姓名", "汉字", "local name", "name in chinese", "chinese name", "native name"), as_text(100)),
        F("no", ("no", "#", "序号", "编号", "№", "n", "nr", "номер п/п", "ნომერი")),
        F("first_name", ("first name", "given name", "given names", "სახელი (first)", "имя (first)", "名")),
        F("last_name", ("last name", "surname", "family name", "გვარი", "фамилия", "姓")),
        F("gender", ("gender", "sex", "სქესი", "пол", "性别"), GENDER, example="F"),
        F("birth_date", ("date of birth", "dob", "birth date", "birthday", "born", "დაბადების თარიღი", "დაბადება",
                         "дата рождения", "出生日期", "生日"), as_date, example=date(1985, 4, 12)),
        F("nationality", ("nationality", "citizenship", "country", "მოქალაქეობა", "ეროვნება", "гражданство", "国籍"), as_text(100), example="China"),
        F("passport_number", ("passport", "passport no", "passport number", "pass no", "passport #", "პასპორტი",
                              "პასპორტის ნომერი", "паспорт", "номер паспорта", "护照", "护照号", "护照号码"), as_text(50), example="E12345678"),
        F("passport_expiry", ("passport expiry", "expiry", "expiry date", "date of expiry", "valid until", "expires",
                              "passport valid", "პასპორტის ვადა", "ვადა", "срок действия", "有效期", "护照有效期"), as_date, example=date(2031, 1, 31)),
        F("room_type", ("room type", "room", "accommodation", "rooming", "room arrangement", "ოთახის ტიპი", "ოთახი", "тип номера", "номер",
                        "房型", "分房", "住房", "房间安排", "分房情况", "房间类型"), ROOM, example="DBL"),
        F("room_note", ("room note", "note for hotel", "bed", "beds", "房间备注", "床型"), as_text(300)),
        F("room_group", ("room no", "room number", "room group", "roommate", "roommates", "sharing", "share with",
                         "ოთახის ნომერი", "№ номера", "房号", "房间", "同住"), as_text(40), example="R1"),
        F("phone", PHONE, as_text(50)),
        F("notes", NOTES, as_text(2000), example="Vegetarian"),
    ],
    "tours": [
        F("code", ("code", "tour code", "tour", "group", "group code", "ტურის კოდი", "კოდი", "ტური", "код", "团号"), as_text(40), example="HM-1106"),
        F("route", ("route", "series", "program", "programme", "მარშრუტი", "სერია", "маршрут", "线路"), as_text(20), example="HM"),
        F("start_date", ("start", "start date", "from", "arrival", "date from", "დაწყება", "დაწყების თარიღი", "ჩამოსვლა",
                         "начало", "заезд", "开始", "抵达"), as_date, required=True, example=date(2026, 11, 6)),
        F("end_date", ("end", "end date", "to", "departure", "date to", "დასრულება", "დასრულების თარიღი", "გამგზავრება",
                       "конец", "выезд", "结束", "离开"), as_date),
        F("client", ("client", "agency", "partner", "customer", "კლიენტი", "სააგენტო", "клиент", "агентство", "客户", "旅行社"), as_text(200), example="Orient Holidays"),
        F("pax", ("pax", "people", "persons", "tourists", "group size", "ტურისტები", "რაოდენობა", "кол-во", "человек", "人数"), as_int, example=24),
        F("rooms", ("rooms", "rooming", "ოთახები", "номера", "房间"), as_text(200), example="12 DBL"),
        F("status", ("status", "სტატუსი", "статус", "状态"), STATUS, example="confirmed"),
        F("name", ("name", "title", "tour name", "სახელი", "название"), as_text(200)),
        F("notes", NOTES, as_text(5000)),
    ],
    "vendors": [
        F("name", NAME, as_text(200), required=True, example="Old Town Hotel"),
        F("type", ("type", "category", "kind", "ტიპი", "კატეგორია", "тип", "类型"), VENDOR_TYPE, required=True, example="hotel"),
        F("city", ("city", "town", "location", "ქალაქი", "город", "城市"), as_text(100), example="Tbilisi"),
        F("contact_person", ("contact", "contact person", "manager", "საკონტაქტო პირი", "контакт", "联系人"), as_text(200)),
        F("phone", PHONE, as_text(50), example="+995 555 000 000"),
        F("phone2", ("phone 2", "phone2", "second phone", "მეორე ტელეფონი", "телефон 2"), as_text(50)),
        F("email", EMAIL, as_text(255)),
        F("payment_days", ("payment days", "days", "pay days", "გადახდის დღეები", "дней оплаты"), as_int, example=7),
        F("payment_timing", ("payment timing", "timing", "pay", "გადახდა", "оплата"), TIMING, example="after"),
        F("currency", ("currency", "ვალუტა", "валюта", "币种"), as_currency),
        F("notes", NOTES, as_text(5000)),
    ],
    "clients": [
        F("name", NAME, as_text(200), required=True, example="Orient Holidays"),
        F("type", ("type", "ტიპი", "тип"), CLIENT_TYPE, example="agency"),
        F("country", ("country", "ქვეყანა", "страна", "国家"), as_text(100), example="China"),
        F("contact_person", ("contact", "contact person", "manager", "საკონტაქტო პირი", "контакт", "联系人"), as_text(200)),
        F("email", EMAIL, as_text(255)),
        F("phone", PHONE, as_text(50)),
        F("notes", NOTES, as_text(5000)),
    ],
}


def _alias_table(spec: list[F]) -> dict[str, str]:
    # Earlier fields win, so a bare "name" means full_name before first_name.
    table: dict[str, str] = {}
    for f in spec:
        for a in (f.key, *f.aliases):
            table.setdefault(norm_header(a), f.key)
    return table


@dataclass
class Parsed:
    columns: dict[str, str | None]
    unread: int = 0  # headed columns not matched to a field
    total: int = 0  # data rows under the header
    records: list[tuple[int, dict]] = field(default_factory=list)
    errors: list[ImportIssue] = field(default_factory=list)
    warnings: list[ImportIssue] = field(default_factory=list)


def parse(entity: str, table: list[tuple[int, list]], forced: dict | None = None) -> Parsed:
    """`forced`: {"header_row": n, "columns": {index: field}, "leader": "TL"} - which column is which, as AI read it."""
    spec = SPECS[entity]
    aliases = _alias_table(spec)
    by_key = {f.key: f for f in spec}

    # Header = the first of the top rows that names at least two known columns.
    header_idx = 0
    for i, (_, row) in enumerate(table[:10]):
        if len({aliases[n] for n in map(norm_header, row) if n in aliases}) >= 2:
            header_idx = i
            break
    if forced:
        header_idx = next((i for i, (n, _) in enumerate(table) if n == forced.get("header_row")), -1)
    _, header = table[header_idx] if header_idx >= 0 else (0, [])
    col_map: dict[int, str] = {}
    columns: dict[str, str | None] = {}
    for idx, h in enumerate([] if forced else header):
        if _blank(h) is None:
            # A column without a header right of one that has it: "英文名" over SURNAME | Given, or a note
            # next to the room ("needs two beds").
            left = col_map.get(idx - 1)
            if entity == "participants" and left == "en_name" and "first_name" not in col_map.values():
                col_map[idx] = "first_name"
            elif entity == "participants" and left == "room_type" and "room_note" not in col_map.values():
                col_map[idx] = "room_note"
            continue
        key = aliases.get(norm_header(h))
        if key in col_map.values():
            key = None  # second column for the same field: ignore it
        columns[str(h).strip()] = key
        if key:
            col_map[idx] = key

    if forced:
        col_map = {int(i): k for i, k in forced["columns"].items() if k in by_key}
        columns = {(str(header[i]).strip() if i < len(header) and _blank(header[i]) is not None else _col_letter(i)): k
                   for i, k in col_map.items()}
    leader_words = LEADER_WORDS | ({norm_header(forced["leader"])} if forced and forced.get("leader") else set())
    result = Parsed(columns=columns)
    missing = [f.key for f in spec if f.required and f.key not in col_map.values()]
    if entity == "participants" and not {"full_name", "first_name", "last_name", "en_name", "local_name"} & set(col_map.values()):
        missing.append("full_name")
    if missing:
        result.errors.append(ImportIssue(row=table[header_idx][0], message=f"Missing column(s): {', '.join(missing)}"))
        return result

    body = table[header_idx + 1:]
    result.unread = sum(1 for k in columns.values() if k is None)
    if len(body) > MAX_ROWS:
        raise ImportFileError(f"At most {MAX_ROWS} rows per import")
    for rownum, row in body:
        rec: dict = {}
        problems = []
        cells = {idx: (_blank(row[idx]) if idx < len(row) else None) for idx in col_map}
        if all(v is None for i, v in cells.items() if col_map[i] != "no"):
            continue  # nothing in the recognised columns (e.g. only a row number)
        result.total += 1
        for idx, key in col_map.items():
            raw = cells[idx]
            if raw is None:
                continue
            try:
                rec[key] = by_key[key].parse(raw)
            except ValueError as e:
                problems.append(f"{key}: {e}")
        if entity == "participants":
            en = rec.pop("en_name", "")
            first, last = rec.pop("first_name", ""), rec.pop("last_name", "") or (en if "first_name" in col_map.values() else "")
            if not rec.get("full_name") and (first or last):
                # Passport order: SURNAME Given names
                rec["full_name"] = f"{str(last).upper()} {first}".strip()
            if not rec.get("full_name") and en:
                rec["full_name"] = en
            if not rec.get("full_name") and rec.get("local_name"):
                rec["full_name"] = rec["local_name"]
            if norm_header(rec.pop("no", "")) in leader_words:
                rec["leader"] = True
            if not rec.get("full_name"):
                problems.append("full_name: missing")
        for f in spec:
            if f.required and f.key not in rec:
                problems.append(f"{f.key}: missing")
        if problems:
            result.errors.append(ImportIssue(row=rownum, message="; ".join(problems)))
        else:
            result.records.append((rownum, rec))
    if entity == "participants":
        room_col = next((i for i, k in col_map.items() if k == "room_type"), None)
        _rooms(result, table, room_col)
        _passports(result, table)
    return result


CAPACITY = {"SGL": 1, "DBL": 2, "TWN": 2, "TRP": 3, "QUAD": 4}


def _rooms(result: "Parsed", table, room_col: int | None) -> None:
    """Who shares a room, as such lists show it: one room cell merged over the roommates' rows, or a room
    type on the first of them and the next rows empty (a twin takes the next person, a triple two more).
    Each room gets a label (R1, R2…); a list's own room numbers win."""
    merged = getattr(table, "merged", {}) or {}
    n, cur = 0, None          # cur: [label, places left, merged top row, room type]
    for rownum, rec in result.records:
        if rec.get("room_group"):
            cur = None
            continue
        block = merged.get((rownum, room_col)) if room_col is not None else None
        if cur and block and block[0] == cur[2] and rownum != block[0]:
            rec["room_group"], rec["room_type"] = cur[0], rec.get("room_type") or cur[3]
        elif rec.get("room_type"):
            n += 1
            places = block[1] if block and block[0] == rownum and block[1] > 1 else CAPACITY.get(rec["room_type"], 1)
            cur = [f"R{n}", places - 1, block[0] if block else None, rec["room_type"]]
            rec["room_group"] = cur[0]
        elif cur and cur[1] > 0 and room_col is not None:
            rec["room_group"], rec["room_type"] = cur[0], cur[3]
            cur[1] -= 1
        else:
            cur = None
            continue
        if cur and cur[2] is not None and block and block[0] == cur[2] and rownum != block[0]:
            cur[1] -= 1
    # A room note goes to everyone in the room.
    notes: dict[str, str] = {}
    for _, rec in result.records:
        if rec.get("room_group") and rec.get("room_note"):
            notes.setdefault(rec["room_group"], rec["room_note"])
    for _, rec in result.records:
        if rec.get("room_group") in notes and not rec.get("room_note"):
            rec["room_note"] = notes[rec["room_group"]]


PASSPORT_NO = re.compile(r"^[A-Z]{1,2}\d{6,9}$")


def _passports(result: "Parsed", table) -> None:
    """Passports listed on another sheet (name, passport number, dates): matched to the people by name."""
    people: dict[str, dict] = {}
    for _, rec in result.records:
        for k in ("local_name", "full_name"):
            if rec.get(k):
                people.setdefault(norm_header(rec[k]), rec)
    today = date.today()
    for sheet in getattr(table, "others", []) or []:
        for _, row in sheet:
            cells = [_blank(v) for v in row]
            rec = next((people[norm_header(c)] for c in cells if isinstance(c, str) and norm_header(c) in people), None)
            if rec is None:
                continue
            for c in cells:
                if c is None:
                    continue
                token = str(c).replace(" ", "").upper()
                if isinstance(c, str) and PASSPORT_NO.match(token) and not rec.get("passport_number"):
                    rec["passport_number"] = token
                    continue
                try:
                    d = as_date(c) if isinstance(c, (date, datetime)) or re.match(r"^\d{4}[-./]\d{1,2}[-./]\d{1,2}", str(c)) else None
                except ValueError:
                    d = None
                if d is not None:
                    if d > today and not rec.get("passport_expiry"):
                        rec["passport_expiry"] = d
                    elif d < today and not rec.get("birth_date"):
                        rec["birth_date"] = d
                    continue
                if isinstance(c, str) and not rec.get("gender"):
                    try:
                        rec["gender"] = GENDER(c)
                    except ValueError:
                        pass


# --- per-entity checks against the database and saving ---------------------------------

def _ci(col):
    return func.lower(col)


def _col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _shape(v) -> str:
    """A cell for the AI to read the column from, without the person: numbers (passports, phones, dates) as 9."""
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return "9999-99-99"
    return re.sub(r"\d", "9", str(v).strip())[:40]


def ai_columns(ws, table) -> dict | None:
    """Which column is which, read by AI from the first rows (numbers masked) - for lists the usual reading
    doesn't understand. None without an AI key or when the answer isn't usable."""
    from .ai import keys
    if not keys.key_for(ws):
        return None
    sample = "\n".join(f"row {n}: " + " | ".join(f"{_col_letter(i)}={_shape(v)}" for i, v in enumerate(row) if _blank(v) is not None)
                       for n, row in table[:14])
    others = "\n".join(f"other sheet row: " + " | ".join(_shape(v) for v in row if _blank(v) is not None)
                        for sheet in (getattr(table, "others", []) or [])[:1] for _, row in sheet[:3])
    fields = [f.key for f in SPECS["participants"]]
    prompt = (
        "These are the first rows of a tour group's passenger list, sent by a travel agency (any language). Numbers are "
        "masked as 9. Say which column holds what. Answer with only JSON: {\"header_row\": <row number of the header, "
        "or the row before the first person if there is none>, \"columns\": {\"<column letter>\": \"<field>\"}, "
        "\"leader\": \"<the text that marks the tour leader in a column, or empty>\"}. Fields: " + ", ".join(fields) + ". "
        "Use en_name for a Latin passport name in one column; last_name and first_name when it is split in two; "
        "local_name for a name in Chinese or another script; room_type for the room (twin, double, 标间…); room_note for "
        "a note about the room; no for the row number; leave out columns that are none of these.\n\n" + sample
        + ("\n\n" + others if others else ""))
    try:
        from .programme import make_client
        from .config import settings
        msg = make_client(ws).messages.create(model=settings.ai_translate_model, max_tokens=1500,
                                              messages=[{"role": "user", "content": prompt}])
        text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
        got = __import__("json").loads(text[text.index("{"): text.rindex("}") + 1])
        cols = {}
        for letter, key in (got.get("columns") or {}).items():
            idx = 0
            for ch in str(letter).strip().upper():
                idx = idx * 26 + (ord(ch) - 64)
            if idx > 0 and key in fields:
                cols[idx - 1] = key
        return {"header_row": int(got.get("header_row") or 0), "columns": cols, "leader": str(got.get("leader") or "")} if cols else None
    except Exception:          # noqa: BLE001 - no AI: the usual reading stands
        return None


def run_import(db: Session, ctx: WorkspaceCtx, entity: str, table: list[tuple[int, list]], commit: bool,
               tour: Tour | None = None, replace: bool = False) -> ImportResult:
    p = parse(entity, table)
    by_ai = False
    # A list the usual reading doesn't understand (no name column, or several unknown columns): AI reads the columns.
    if entity == "participants" and (not p.records or p.unread >= 2 or any(e.message.startswith("Missing column") for e in p.errors)):
        forced = ai_columns(ctx.workspace, table)
        if forced:
            q = parse(entity, table, forced)
            if len(q.records) >= len(p.records) and q.records:
                p, by_ai = q, True
    created = 0
    preview: list[dict] = []
    valid = 0
    if p.records:
        handler = {"participants": _participants, "tours": _tours, "vendors": _vendors, "clients": _clients}[entity]
        for rownum, rec, obj in handler(db, ctx, p, tour, replace if commit else False):
            valid += 1
            if len(preview) < 200:
                preview.append({"row": rownum, **{k: (v.isoformat() if isinstance(v, date) else v) for k, v in rec.items()}})
            if commit and obj is not None:
                db.add(obj)
                created += 1
        if commit:
            db.commit()
        else:
            db.rollback()
    p.errors.sort(key=lambda e: e.row)
    return ImportResult(total=p.total, valid=valid, created=created, committed=commit, columns=p.columns,
                        errors=p.errors, warnings=p.warnings, preview=preview, by_ai=by_ai)


def _participants(db, ctx, p: Parsed, tour: Tour, replace: bool):
    if replace:
        db.execute(delete(Participant).where(Participant.tour_id == tour.id))
    pos = next_position(db, tour.id)
    limit = tour.end_date + timedelta(days=183)
    seen: set[str] = set()
    leaders = sum(1 for _, rec in p.records if rec.get("leader"))
    if leaders:
        # "领队" in the list: the group's free tour leader (counted in the tour's leaders, eats at the staff table).
        p.warnings.append(ImportIssue(row=0, message=f"{leaders} tour leader(s) found: counted as the tour's free tour leaders"))
        if not tour.leaders:
            tour.leaders = leaders
    for rownum, rec in p.records:
        if rec.pop("leader", False):
            rec["notes"] = ("Tour leader. " + rec.get("notes", "")).strip()
        pn = rec.get("passport_number", "").replace(" ", "").upper()
        if pn:
            if pn in seen:
                p.warnings.append(ImportIssue(row=rownum, message=f"passport {pn} appears twice in the file"))
            seen.add(pn)
        if rec.get("passport_expiry") and rec["passport_expiry"] < limit:
            p.warnings.append(ImportIssue(row=rownum, message="passport expires within 6 months of the tour end"))
        yield rownum, rec, Participant(tour_id=tour.id, position=pos, **rec)
        pos += 1


def _tours(db, ctx, p: Parsed, tour, replace):
    series = {s.code.upper(): s for s in db.scalars(select(Series).where(Series.workspace_id == ctx.workspace.id))}
    clients = {c.name.lower(): c for c in db.scalars(select(Client).where(Client.workspace_id == ctx.workspace.id))}
    codes_in_file: set[str] = set()
    for rownum, rec in p.records:
        s = None
        if rec.get("route"):
            s = series.get(rec["route"].upper())
            if not s:
                p.errors.append(ImportIssue(row=rownum, message=f"route '{rec['route']}' does not exist"))
                continue
        start = rec["start_date"]
        end = rec.get("end_date") or (start + timedelta(days=len(s.days) - 1) if s and s.days else None)
        if end is None:
            p.errors.append(ImportIssue(row=rownum, message="end_date is required when there is no route"))
            continue
        if end < start or (end - start).days + 1 > MAX_TOUR_DAYS:
            p.errors.append(ImportIssue(row=rownum, message="end date must be after start (max 60 days)"))
            continue
        code = (rec.get("code") or "").upper() or _generate_code(db, ctx, s, start)
        while not rec.get("code") and code in codes_in_file:  # two generated codes in one file
            code += "B"
        if code in codes_in_file or _code_taken(db, ctx, code):
            p.errors.append(ImportIssue(row=rownum, message=f"tour code {code} already exists"))
            continue
        codes_in_file.add(code)
        client = None
        if rec.get("client"):
            client = clients.get(rec["client"].lower())
            if not client and not ctx.can("finance", "sales"):
                p.errors.append(ImportIssue(row=rownum, message=f"unknown client '{rec['client']}' — ask a manager to add it"))
                continue
            if not client:
                client = Client(workspace_id=ctx.workspace.id, name=rec["client"])
                clients[rec["client"].lower()] = client
                db.add(client)
                p.warnings.append(ImportIssue(row=rownum, message=f"new client '{rec['client']}' will be created"))
        t = Tour(workspace_id=ctx.workspace.id, code=code, series_id=s.id if s else None, client=client,
                 start_date=start, end_date=end, status=rec.get("status", "confirmed"), pax=min(max(rec.get("pax", 0) or 0, 0), 10000),
                 rooms=rec.get("rooms", ""), name=rec.get("name", ""), notes=rec.get("notes", ""))
        sync_days(t, s.days if s else None)
        yield rownum, {**rec, "code": code, "end_date": end}, t


def _vendors(db, ctx, p: Parsed, tour, replace):
    existing = {(n.lower(), t) for n, t in db.execute(
        select(Vendor.name, Vendor.type).where(Vendor.workspace_id == ctx.workspace.id))}
    for rownum, rec in p.records:
        key = (rec["name"].lower(), rec["type"])
        if key in existing:
            p.errors.append(ImportIssue(row=rownum, message=f"{rec['type']} '{rec['name']}' already exists"))
            continue
        existing.add(key)
        yield rownum, rec, Vendor(workspace_id=ctx.workspace.id, **rec)


def _clients(db, ctx, p: Parsed, tour, replace):
    existing = {n.lower() for n in db.scalars(select(Client.name).where(Client.workspace_id == ctx.workspace.id))}
    for rownum, rec in p.records:
        if rec["name"].lower() in existing:
            p.errors.append(ImportIssue(row=rownum, message=f"client '{rec['name']}' already exists"))
            continue
        existing.add(rec["name"].lower())
        yield rownum, rec, Client(workspace_id=ctx.workspace.id, **rec)


# --- templates -------------------------------------------------------------------------

def template(entity: str) -> bytes:
    spec = [f for f in SPECS[entity] if f.key not in ("first_name", "last_name")]
    wb = Workbook()
    ws = wb.active
    ws.title = entity.capitalize()
    ws.append([f.key for f in spec])
    ws.append([f.example if f.example != "" else None for f in spec])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F46E5")
    for c in ws[2]:
        if isinstance(c.value, date):
            c.number_format = "YYYY-MM-DD"
    for i, f in enumerate(spec):
        ws.column_dimensions[chr(65 + i)].width = max(14, len(f.key) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
