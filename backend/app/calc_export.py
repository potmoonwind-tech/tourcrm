"""The whole calculation as an Excel workbook, the way agencies keep it: prices for every group size,
every line, the hotels night by night, one group size worked through from cost to price, and the
inputs, formulas (in words too) and exchange rates. Values, not live formulas, so the file always
matches the app; each formula sits next to its line as text and in words."""
import io
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session


from .models import Vendor
from .quote_engine import breakdown, compute, following, in_season, with_transport
from .quote_schemas import CalcStyle, QuoteIn
from .travel import journey
from .xlsx_safe import defuse

BRAND = "0073EA"
HEAD = PatternFill("solid", fgColor=BRAND)
SOFT = PatternFill("solid", fgColor="EEF4FF")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
MONEY = "#,##0.00"
THIN = Border(bottom=Side(style="thin", color="D0D4E4"))

T = {
    "en": {
        "prices": "Prices", "lines": "Lines", "hotels": "Hotels", "inputs": "Inputs & formulas",
        "agency": "Agency", "calc": "Calculation", "version": "Version", "exported": "Exported", "currency": "Currency",
        "season": "Season", "base": "base prices", "valid": "Valid", "group_at": "Lines at group size",
        "c_group": "Group", "c_leaders": "Free", "c_cost": "Cost per person", "c_profit": "Profit per person",
        "c_dbl": "Price, double", "c_sgl": "Price, single", "c_sup": "Single supplement", "c_trp": "Price, triple",
        "c_quad": "Price, room for 4", "c_gcost": "Group cost", "c_grev": "Group revenue", "c_gprofit": "Group profit",
        "c_vehicle": "Vehicle", "c_day": "Day", "c_kind": "Kind", "c_line": "Line", "c_supplier": "Supplier",
        "c_unit": "Unit price", "c_cur": "Currency", "c_charged": "Charged", "c_who": "Who", "c_qtyf": "How many (formula)",
        "c_qty": "How many", "c_mult": "× people / vehicles", "c_total": "For the group", "c_pp": "Per paying person",
        "c_in": "In the price", "c_words": "Formula in words", "c_notes": "Notes", "c_night": "Night", "c_city": "City",
        "c_hotel": "Hotel", "c_rdbl": "Double room", "c_rsgl": "Single room", "c_rtrp": "Triple room", "c_rquad": "Room for 4",
        "c_bed": "Bed per person (double)", "c_staff": "Driver & guide", "c_staffcost": "Driver & guide cost",
        "c_name": "Name", "c_value": "Value", "c_note": "Note", "c_what": "What", "c_formula": "Formula", "c_reads": "Reads as",
        "c_type": "Type of person", "c_role": "Role", "c_count": "How many", "c_own": "Own costs per person",
        "c_shared": "Share of shared costs", "c_for": "What they cost the group",
        "per_group": "once for the group", "per_person": "per person", "per_vehicle": "per vehicle",
        "yes": "yes", "no": "no", "optional": "optional extra",
        "w_tourists": "who pays", "w_persons": "who pays + free", "w_everyone": "everyone", "w_staff": "driver & guide", "w_crew": "leader + staff",
        "s_lines": "All lines, for the group", "s_hotel_dbl": "Hotels, bed per person (double)", "s_hotel_sgl": "Hotels, single room per person",
        "s_leader": "Rooms of people travelling free", "s_staffroom": "Driver & guide rooms", "s_shared": "Shared costs per paying person",
        "s_cost_dbl": "Cost per person, double", "s_cost_sgl": "Cost per person, single", "s_cost": "Cost",
        "s_price_dbl": "Price per person, double", "s_price_sgl": "Price per person, single", "s_steps": "From cost to price",
        "s_people": "Who travels", "s_after": "price after it",
        "price": "price", "quantity": "how many", "comes_with": "comes with", "step": "pricing step",
        "rates": "Exchange rates", "rate_line": "1 {base} = {rate} {cur}", "no_formulas": "No formulas",
        "stay_none": "not counted here", "stay_same": "same hotel", "stay_own": "other hotel", "stay_free": "free from the hotel",
        "stay_allowance": "allowance (they find a hotel)",
        "nm": {"days": "days of the trip", "nights": "hotel nights", "group": "the group", "tourists": "people who pay",
               "persons": "people (who pay + free)", "everyone": "everyone", "leaders": "people travelling free",
               "staff": "driver & guide", "crew": "leader + driver & guide", "cost_total": "all costs"},
        "cost_of": "the cost of {what}", "pct": "{p}% of {of}", "per": "one for every {n} {who}, rounded up",
        "per_more": "one for every {n} {who} (+{more}), rounded up", "if": "if {name} {op} {val}: {yes}, otherwise {no}",
        "least": "{k} × {who}, but at least {min}", "minus": "{k} × {what}, minus {less}",
        "ops": {"==": "is", ">": "is more than", ">=": "is at least", "<": "is less than", "<=": "is at most"},
        "words": {"if": "if", "else": "otherwise", "ceil": "round up", "floor": "round down", "max": "the bigger of",
                  "min": "the smaller of", "round": "round", "abs": "without sign", "and": "and", "or": "or"},
        "kinds": {"hotel": "hotels", "meal": "meals", "transport": "transport", "guide": "guide", "ticket": "tickets",
                  "activity": "activities", "tips": "tips", "border": "border", "other": "other"},
    },
    "ka": {
        "prices": "ფასები", "lines": "ხაზები", "hotels": "სასტუმროები", "inputs": "მონაცემები და ფორმულები",
        "agency": "სააგენტო", "calc": "კალკულაცია", "version": "ვერსია", "exported": "ექსპორტი", "currency": "ვალუტა",
        "season": "სეზონი", "base": "საბაზო ფასები", "valid": "მოქმედებს", "group_at": "ხაზები ჯგუფისთვის",
        "c_group": "ჯგუფი", "c_leaders": "უფასო", "c_cost": "თვითღირებულება ერთზე", "c_profit": "მოგება ერთზე",
        "c_dbl": "ფასი, ორადგილიანი", "c_sgl": "ფასი, ერთადგილიანი", "c_sup": "სინგლის დანამატი", "c_trp": "ფასი, სამადგილიანი",
        "c_quad": "ფასი, ოთხადგილიანი", "c_gcost": "ჯგუფის ხარჯი", "c_grev": "ჯგუფის შემოსავალი", "c_gprofit": "ჯგუფის მოგება",
        "c_vehicle": "ტრანსპორტი", "c_day": "დღე", "c_kind": "სახეობა", "c_line": "ხაზი", "c_supplier": "მომწოდებელი",
        "c_unit": "ერთ. ფასი", "c_cur": "ვალუტა", "c_charged": "როგორ ითვლება", "c_who": "ვინ", "c_qtyf": "რამდენი (ფორმულა)",
        "c_qty": "რამდენი", "c_mult": "× ადამიანი / მანქანა", "c_total": "ჯგუფზე", "c_pp": "ერთ გადამხდელზე",
        "c_in": "ფასშია", "c_words": "ფორმულა სიტყვებით", "c_notes": "შენიშვნა", "c_night": "ღამე", "c_city": "ქალაქი",
        "c_hotel": "სასტუმრო", "c_rdbl": "ორადგილიანი", "c_rsgl": "ერთადგილიანი", "c_rtrp": "სამადგილიანი", "c_rquad": "ოთხადგილიანი",
        "c_bed": "საწოლი ერთზე (ორადგ.)", "c_staff": "მძღოლი და გიდი", "c_staffcost": "მძღოლისა და გიდის ხარჯი",
        "c_name": "სახელი", "c_value": "მნიშვნელობა", "c_note": "შენიშვნა", "c_what": "რა", "c_formula": "ფორმულა", "c_reads": "იკითხება ასე",
        "c_type": "ადამიანის ტიპი", "c_role": "როლი", "c_count": "რამდენი", "c_own": "საკუთარი ხარჯი ერთზე",
        "c_shared": "საერთო ხარჯის წილი", "c_for": "რა უჯდება ჯგუფს",
        "per_group": "ერთხელ ჯგუფზე", "per_person": "ერთ ადამიანზე", "per_vehicle": "მანქანაზე",
        "yes": "კი", "no": "არა", "optional": "დამატებითი",
        "w_tourists": "ვინც იხდის", "w_persons": "იხდის + უფასო", "w_everyone": "ყველა", "w_staff": "მძღოლი და გიდი", "w_crew": "ლიდერი + პერსონალი",
        "s_lines": "ყველა ხაზი, ჯგუფზე", "s_hotel_dbl": "სასტუმრო, საწოლი ერთზე (ორადგ.)", "s_hotel_sgl": "სასტუმრო, ერთადგილიანი ერთზე",
        "s_leader": "უფასოდ მოგზაურების ნომრები", "s_staffroom": "მძღოლისა და გიდის ნომრები", "s_shared": "საერთო ხარჯი ერთ გადამხდელზე",
        "s_cost_dbl": "თვითღირებულება ერთზე, ორადგ.", "s_cost_sgl": "თვითღირებულება ერთზე, ერთადგ.", "s_cost": "თვითღირებულება",
        "s_price_dbl": "ფასი ერთზე, ორადგ.", "s_price_sgl": "ფასი ერთზე, ერთადგ.", "s_steps": "თვითღირებულებიდან ფასამდე",
        "s_people": "ვინ მოგზაურობს", "s_after": "ფასი ამის შემდეგ",
        "price": "ფასი", "quantity": "რამდენი", "comes_with": "მოყვება", "step": "ფასის ნაბიჯი",
        "rates": "ვალუტის კურსები", "rate_line": "1 {base} = {rate} {cur}", "no_formulas": "ფორმულები არ არის",
        "stay_none": "აქ არ ითვლება", "stay_same": "იგივე სასტუმრო", "stay_own": "სხვა სასტუმრო", "stay_free": "სასტუმროსგან უფასო",
        "stay_allowance": "თანხა (სასტუმროს თავად პოულობენ)",
        "nm": {"days": "მოგზაურობის დღეები", "nights": "სასტუმროს ღამეები", "group": "ჯგუფი", "tourists": "ვინც იხდის",
               "persons": "ადამიანები (იხდის + უფასო)", "everyone": "ყველა", "leaders": "უფასოდ მოგზაურები",
               "staff": "მძღოლი და გიდი", "crew": "ლიდერი + მძღოლი და გიდი", "cost_total": "ყველა ხარჯი"},
        "cost_of": "ხარჯი: {what}", "pct": "{of} — {p}%", "per": "ერთი ყოველ {n} {who}-ზე, ზემოთ დამრგვალებით",
        "per_more": "ერთი ყოველ {n} {who}-ზე (+{more}), ზემოთ დამრგვალებით", "if": "თუ {name} {op} {val}: {yes}, თუ არა — {no}",
        "least": "{k} × {who}, მაგრამ მინიმუმ {min}", "minus": "{k} × {what}, მინუს {less}",
        "ops": {"==": "არის", ">": "მეტია", ">=": "მინიმუმ", "<": "ნაკლებია", "<=": "მაქსიმუმ"},
        "words": {"if": "თუ", "else": "თუ არა", "ceil": "ზემოთ დამრგვალება", "floor": "ქვემოთ დამრგვალება", "max": "უდიდესი",
                  "min": "უმცირესი", "round": "დამრგვალება", "abs": "ნიშნის გარეშე", "and": "და", "or": "ან"},
        "kinds": {"hotel": "სასტუმრო", "meal": "კვება", "transport": "ტრანსპორტი", "guide": "გიდი", "ticket": "ბილეთები",
                  "activity": "აქტივობები", "tips": "ჩაი", "border": "საზღვარი", "other": "სხვა"},
    },
}

# The engine's names for steps without a label of their own, in Georgian.
STEP_KA = {"Profit": "მოგება", "Markup": "ნამატი", "Per person": "ერთ ადამიანზე", "Per group": "ჯგუფზე", "Commission": "საკომისიო",
           "VAT on margin": "დღგ მარჟაზე", "VAT": "დღგ", "Discount": "ფასდაკლება", "Share of another price": "სხვა ფასის წილი",
           "Rounding": "დამრგვალება", "Formula": "ფორმულა"}
TOKEN = re.compile(r"\s*(\d+(?:\.\d+)?|[A-Za-z_][A-Za-z0-9_]*|==|!=|>=|<=|[-+*/()<>,%])")
SIGN = {"*": "×", "/": "÷", "-": "−", "==": "=", ">=": "≥", "<=": "≤"}


class _Words:
    """A formula in plain words, in the workbook's language (same patterns as the app's builder)."""

    def __init__(self, L: dict, doc: QuoteIn, own_kinds: dict[str, str]):
        self.L, self.doc, self.own = L, doc, own_kinds

    def name(self, n: str) -> str:
        if n in self.L["nm"]:
            return self.L["nm"][n]
        if n.startswith("cost_"):
            k = n[5:]
            return self.L["cost_of"].format(what=self.own.get(k) or self.L["kinds"].get(k, k))
        return self.doc.variable_notes.get(n) or n

    def __call__(self, f: str) -> str:
        f, L, nm = f.strip(), self.L, self.name
        g = lambda x: f"{float(x):g}"
        if m := re.fullmatch(r"cost_([a-z0-9_]+)\s*\*\s*([\d.]+)", f, re.I):
            return L["pct"].format(p=f"{float(m[2]) * 100:g}", of=nm(f"cost_{m[1]}"))
        if m := re.fullmatch(r"ceil\(\(?\s*([a-z_]\w*)(?:\s*\+\s*(\d+))?\s*\)?\s*/\s*([\d.]+)\s*\)", f, re.I):
            return (L["per_more"] if m[2] else L["per"]).format(n=g(m[3]), who=nm(m[1]), more=m[2])
        if m := re.fullmatch(r"([\d.]+)\s+if\s+([a-z_]\w*)\s*(==|>=|<=|>|<)\s*([\d.]+)\s+else\s+([\d.]+)", f, re.I):
            return L["if"].format(name=nm(m[2]), op=L["ops"][m[3]], val=g(m[4]), yes=g(m[1]), no=g(m[5]))
        if m := re.fullmatch(r"max\(\s*([\d.]+)\s*,\s*([a-z_]\w*)\s*\*\s*([\d.]+)\s*\)", f, re.I):
            return L["least"].format(k=g(m[3]), who=nm(m[2]), min=g(m[1]))
        if m := re.fullmatch(r"([a-z_]\w*)\s*\*\s*([\d.]+)\s*-\s*([a-z_]\w*)", f, re.I):
            return L["minus"].format(k=g(m[2]), what=nm(m[1]), less=nm(m[3]))
        out = []
        for x in TOKEN.findall(f):
            out.append(L["words"].get(x) or SIGN.get(x) or (nm(x) if re.match(r"[A-Za-z_]", x) else x))
        return " ".join(out).replace("( ", "(").replace(" )", ")")


def _plain(expr: str) -> bool:
    return bool(re.fullmatch(r"\s*\d+(\.\d+)?\s*", expr or ""))


def _sheet(wb: Workbook, title: str, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = re.sub(r"[\[\]:*?/\\]", "-", title)[:31]
    ws.sheet_view.showGridLines = False
    return ws


def _header(ws, cols: list[str], widths: list[int]):
    ws.append(cols)
    r = ws.max_row
    for c in ws[r]:
        c.font, c.fill = WHITE_BOLD, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 32
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(r + 1, 1)
    return r


def _money(ws, row: int, cols: list[int]):
    for c in cols:
        ws.cell(row, c).number_format = MONEY


def _title(ws, lines: list[tuple[str, object]]):
    for k, v in lines:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(color="676879")
    ws.cell(1, 2).font = Font(bold=True, size=14)
    ws.append([])


def workbook(db: Session, ws_obj, q, doc: QuoteIn, version: int, lang: str, tourists: int, season: str | None) -> bytes:
    L = T["ka" if lang == "ka" else "en"]
    style = CalcStyle.model_validate(ws_obj.calc_style or {})
    own = {k.id: k.name for k in style.kinds}
    kind_name = lambda it: own.get(it.kind or "") or L["kinds"].get(it.category, it.category)
    words = _Words(L, doc, own)
    sn = next((s for s in doc.seasons if s.id == season), None)
    priced = in_season(doc, sn.id) if sn else doc
    result = compute(priced)
    if result.errors:
        raise ValueError("; ".join(result.errors))
    n = min(max(tourists, doc.pax_min), doc.pax_max)
    b = breakdown(priced, n)
    full = following(with_transport(priced))        # the lines as the engine prices them (transport plan included)
    vendors = {v.id: v.name for v in db.scalars(select(Vendor).where(Vendor.id.in_({i.vendor_id for i in full.items if i.vendor_id} | {x.vendor_id for x in doc.nights if x.vendor_id})))}
    cur = doc.currency
    wb = Workbook()
    head = [(L["calc"], doc.name), (L["agency"], ws_obj.name), (L["version"], f"v{version}"), (L["exported"], date.today().isoformat()),
            (L["currency"], cur), (L["season"], sn.name if sn else L["base"])]
    if doc.valid_from or doc.valid_to:
        head.append((L["valid"], f"{doc.valid_from or '…'} – {doc.valid_to or '…'}"))

    # 1 · Prices for every group size
    ws = _sheet(wb, L["prices"], first=True)
    _title(ws, head)
    trp = any(t.price_trp is not None for t in result.tiers)
    quad = any(t.price_quad is not None for t in result.tiers)
    others = [x.name for x in (result.tiers[0].types[1:] if result.tiers else [])]
    curs = [o.currency for o in (result.tiers[0].other if result.tiers else [])]
    cols = [L["c_group"], L["c_leaders"], L["c_cost"], L["c_profit"], L["c_dbl"], L["c_sgl"], L["c_sup"]]
    cols += ([L["c_trp"]] if trp else []) + ([L["c_quad"]] if quad else []) + others + [f"{L['c_dbl']} · {c}" for c in curs]
    cols += [L["c_gcost"], L["c_grev"], L["c_gprofit"], L["c_vehicle"]]
    _header(ws, cols, [10, 8] + [14] * (len(cols) - 3) + [24])
    for t in result.tiers:
        row = [t.label, t.leaders, t.cost_dbl, t.price_dbl - t.cost_dbl, t.price_dbl, t.price_sgl, t.single_supplement]
        row += ([t.price_trp] if trp else []) + ([t.price_quad] if quad else [])
        row += [x.price for x in t.types[1:]] + [o.price_dbl for o in t.other]
        row += [t.group_cost, t.group_revenue, t.group_revenue - t.group_cost, t.vehicle]
        ws.append(row)
        _money(ws, ws.max_row, list(range(3, len(row))))
        if t.tourists == n:
            for c in ws[ws.max_row]:
                c.fill, c.font = SOFT, BOLD

    # 2 · Every line at the chosen size
    ws = _sheet(wb, L["lines"])
    _title(ws, head + [(L["group_at"], b.tier.label)])
    cols = [L["c_day"], L["c_kind"], L["c_line"], L["c_supplier"], L["c_unit"], L["c_cur"], L["c_charged"], L["c_who"],
            L["c_qtyf"], L["c_qty"], L["c_mult"], f"{L['c_total']}, {cur}", f"{L['c_pp']}, {cur}", L["c_in"], L["c_words"], L["c_notes"]]
    _header(ws, cols, [6, 14, 38, 22, 11, 8, 16, 16, 24, 9, 9, 14, 14, 12, 40, 30])
    paying = max(1, sum(x.count for x in b.tier.types) if b.tier.types else n)
    for it, c in zip(full.items, b.items):
        if it.id.startswith(("transport:", "fuel:")) and not c.quantity:
            continue    # a vehicle of the plan this group size doesn't take
        name = f"{it.name} · {journey(it)}" if it.travel and journey(it) else it.name
        words_ = "; ".join(x for x in [
            f"{L['price']}: {words(it.price_formula)}" if it.price_formula else "",
            f"{L['quantity']}: {words(it.quantity)}" if not _plain(it.quantity) else "",
            f"{L['comes_with']}: {next((x.name for x in doc.items if x.id == it.with_line), '—')}" if it.with_line else ""] if x)
        state = L["optional"] if it.optional and not it.included else (L["yes"] if c.included else L["no"])
        ws.append([it.day or "", kind_name(it), name, vendors.get(it.vendor_id or "", ""), c.unit_price, it.currency,
                   L[f"per_{it.basis.replace('per_', '')}" if it.basis != "group" else "per_group"],
                   L.get(f"w_{it.who}", it.who) if it.basis != "group" else "", it.quantity if not _plain(it.quantity) else "",
                   c.quantity, c.multiplier, c.amount, c.amount / paying, state, words_, (it.notes or "")[:300]])
        _money(ws, ws.max_row, [5, 12, 13])
    ws.append([])
    ws.append(["", "", L["s_lines"], "", "", "", "", "", "", "", "", sum(x.amount for x in b.items)])
    ws.cell(ws.max_row, 3).font = ws.cell(ws.max_row, 12).font = BOLD
    _money(ws, ws.max_row, [12])

    # 3 · Hotels night by night
    ws = _sheet(wb, L["hotels"])
    _title(ws, head)
    seasons = doc.seasons
    cols = [L["c_night"], L["c_city"], L["c_hotel"], L["c_cur"], L["c_rdbl"], L["c_rsgl"], L["c_rtrp"], L["c_rquad"],
            f"{L['c_bed']}, {cur}", L["c_staff"], f"{L['c_staffcost']}, {cur}"] + [f"{s.name} · {L['c_rdbl']}" for s in seasons]
    _header(ws, cols, [7, 14, 28, 8, 12, 12, 12, 12, 16, 26, 16] + [16] * len(seasons))
    rows = {r["night"]: r for r in b.hotel_nights}
    for i, nt in enumerate(doc.nights, 1):
        r = rows.get(i, {})
        stay = nt.staff_stay
        staff = L[f"stay_{stay}"] + (f" · {nt.staff_price:f} {nt.staff_currency or nt.currency}" if stay in ("own", "allowance") and nt.staff_price is not None else "")
        ws.append([i, nt.city, nt.hotel_name or vendors.get(nt.vendor_id or "", ""), nt.currency, nt.price_dbl, nt.price_sgl, nt.price_trp,
                   nt.price_quad, r.get("dbl_pp"), staff, r.get("staff")] +
                  [(nt.season_prices.get(s.id).price_dbl if nt.season_prices.get(s.id) else None) for s in seasons])
        _money(ws, ws.max_row, [5, 6, 7, 8, 9, 11] + list(range(12, 12 + len(seasons))))

    # 4 · One group size, from cost to price (like a "19+1" sheet)
    ws = _sheet(wb, b.tier.label)
    _title(ws, head)
    st = b.steps
    pairs = [(L["s_lines"], st.lines_total), (L["s_leader"], st.leader_room), (L["s_staffroom"], st.staff_room),
             (L["s_shared"], st.shared_pp), (L["s_hotel_dbl"], st.hotel_dbl_pp), (L["s_hotel_sgl"], st.hotel_sgl_pp),
             (L["s_cost_dbl"], b.tier.cost_dbl), (L["s_cost_sgl"], b.tier.cost_sgl)]
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 40, 16, 16
    for k, v in pairs:
        ws.append([k, v])
        ws.cell(ws.max_row, 2).number_format = MONEY
        ws.cell(ws.max_row, 1).border = ws.cell(ws.max_row, 2).border = THIN
    ws.append([])
    ws.append([L["s_steps"], cur, L["s_after"]])
    for c in ws[ws.max_row]:
        c.font, c.fill = WHITE_BOLD, HEAD
    ws.append([L["s_cost"], b.tier.cost_dbl, b.tier.cost_dbl])
    for x in st.recipe:
        ws.append([STEP_KA.get(x.label, x.label) if lang == "ka" else x.label, x.amount, x.price])
        _money(ws, ws.max_row, [2, 3])
    for k, v in [(L["s_price_dbl"], b.tier.price_dbl), (L["s_price_sgl"], b.tier.price_sgl), (L["c_gcost"], b.tier.group_cost),
                 (L["c_grev"], b.tier.group_revenue), (L["c_gprofit"], b.tier.group_revenue - b.tier.group_cost), (L["c_vehicle"], b.tier.vehicle)]:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = BOLD
        if isinstance(v, float):
            ws.cell(ws.max_row, 2).number_format = MONEY
    ws.append([])
    ws.append([L["s_people"]])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    ws.append([L["c_type"], L["c_count"], L["c_own"], L["c_shared"], L["c_for"]])
    for c in ws[ws.max_row]:
        c.font, c.fill = WHITE_BOLD, HEAD
    for p in b.people:
        ws.append([p.name, p.count, p.own or None, p.shared or None, p.for_others or None])
        _money(ws, ws.max_row, [3, 4, 5])
    for col in "DE":
        ws.column_dimensions[col].width = 18

    # 5 · Inputs, formulas, exchange rates
    ws = _sheet(wb, L["inputs"])
    _title(ws, head)
    _header(ws, [L["c_name"], L["c_value"], L["c_note"]], [28, 14, 50])
    for k, v in doc.variables.items():
        ws.append([k, float(v), doc.variable_notes.get(k, "")])
    ws.append([])
    ws.append([L["c_line"], L["c_what"], L["c_formula"], L["c_reads"]])
    for c in ws[ws.max_row]:
        c.font, c.fill = WHITE_BOLD, HEAD
    ws.column_dimensions["D"].width = 50
    found = 0
    for it in doc.items:
        for what, expr in [(L["price"], it.price_formula), (L["quantity"], "" if _plain(it.quantity) else it.quantity)]:
            if expr:
                ws.append([it.name, what, expr, words(expr)])
                found += 1
        if it.with_line:
            ws.append([it.name, L["comes_with"], "", next((x.name for x in doc.items if x.id == it.with_line), "—")])
            found += 1
    for s in doc.recipe:
        if s.kind == "formula" and s.formula:
            ws.append([s.label or "Formula", L["step"], s.formula, words(s.formula)])
            found += 1
    if not found:
        ws.append([L["no_formulas"]])
    ws.append([])
    ws.append([L["rates"]])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    for c, rate in doc.fx.items():
        ws.append([L["rate_line"].format(base=cur, rate=f"{rate:f}".rstrip("0").rstrip("."), cur=c)])

    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
    defuse(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()



