"""Who works a tour (staff assignments) and who travels on it (participants).

Participant rows hold passport data, so they're office-only (staff and up),
like the rest of the tour API; guides will get a reduced view in their portal.
"""
import io
import math
import re
from collections import Counter, defaultdict
from datetime import timedelta

from fastapi import APIRouter, Depends, HTTPException, Response, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from . import staff_change
from ..db import get_db
from ..xlsx_safe import defuse
from ..deps import WorkspaceCtx, require_role
from ..models import Participant, StaffAssignment, Tour, Vendor
from ..locks import TourLocked
from ..schemas import (
    ParticipantIn, ParticipantOut, ParticipantUpdate, RoomingSummary, StaffIn, StaffOut, StaffUpdate,
)
from ..scoping import check_refs, get_owned

router = APIRouter(prefix="/api/workspaces/{workspace_id}/tours/{tour_id}", tags=["people"])

office = require_role("staff")

ROOM_CAPACITY = {"SGL": 1, "DBL": 2, "TWN": 2, "TRP": 3, "QUAD": 4}
PASSPORT_MARGIN = timedelta(days=183)  # most countries want 6 months' validity


def _tour(db: Session, ctx: WorkspaceCtx, tour_id: str) -> Tour:
    return get_owned(db, ctx, Tour, tour_id)


def _tour_days(t: Tour) -> int:
    return (t.end_date - t.start_date).days + 1


# --- staff -------------------------------------------------------------------------

def _span(t: Tour, start: int, end: int | None) -> str:
    return staff_change.days_text(start, end or _tour_days(t))


def _check_range(t: Tour, start: int, end: int | None) -> None:
    n = _tour_days(t)
    if start > n or (end is not None and (end > n or end < start)):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, f"Days must be within 1–{n}, start before end")


@router.get("/staff", response_model=list[StaffOut])
def list_staff(tour_id: str, ctx: WorkspaceCtx = Depends(office), db: Session = Depends(get_db)):
    return _tour(db, ctx, tour_id).staff


@router.post("/staff", response_model=StaffOut, status_code=status.HTTP_201_CREATED)
def add_staff(tour_id: str, body: StaffIn, ctx: WorkspaceCtx = Depends(office), db: Session = Depends(get_db)):
    t = _tour(db, ctx, tour_id)
    check_refs(db, ctx, {"vendor_id": (Vendor, body.vendor_id)})
    _check_range(t, body.start_day, body.end_day)
    a = StaffAssignment(tour_id=t.id, **body.model_dump(exclude_none=True))
    db.add(a)
    v = db.get(Vendor, body.vendor_id)
    staff_change.record(db, t, f"{_span(t, a.start_day, a.end_day)}: {a.role.replace('_', ' ')} {v.name if v else ''} added", ctx.user.id)
    db.commit()
    db.refresh(a)
    return a


def _assignment(db: Session, t: Tour, staff_id: str) -> StaffAssignment:
    a = db.scalar(select(StaffAssignment).where(StaffAssignment.id == staff_id, StaffAssignment.tour_id == t.id))
    if not a:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Assignment not found")
    return a


@router.patch("/staff/{staff_id}", response_model=StaffOut)
def update_staff(tour_id: str, staff_id: str, body: StaffUpdate, ctx: WorkspaceCtx = Depends(office),
                 db: Session = Depends(get_db)):
    t = _tour(db, ctx, tour_id)
    a = _assignment(db, t, staff_id)
    values = body.model_dump(exclude_unset=True)
    if values.get("vendor_id"):
        check_refs(db, ctx, {"vendor_id": (Vendor, values["vendor_id"])})
    start = values.get("start_day") or a.start_day
    end = values["end_day"] if "end_day" in values else a.end_day
    _check_range(t, start, end)
    was = (a.vendor.name, a.start_day, a.end_day)
    for k, v in values.items():
        if v is not None or k == "end_day":
            setattr(a, k, v)
    db.flush()
    db.refresh(a, ["vendor"])
    if was != (a.vendor.name, a.start_day, a.end_day):
        staff_change.record(db, t, f"{a.role.replace('_', ' ').capitalize()}: {was[0]} {_span(t, was[1], was[2]).lower()} → "
                                   f"{a.vendor.name} {_span(t, a.start_day, a.end_day).lower()}", ctx.user.id)
    db.commit()
    db.refresh(a)
    return a


@router.delete("/staff/{staff_id}", status_code=status.HTTP_204_NO_CONTENT)
def remove_staff(tour_id: str, staff_id: str, ctx: WorkspaceCtx = Depends(office), db: Session = Depends(get_db)):
    t = _tour(db, ctx, tour_id)
    a = _assignment(db, t, staff_id)
    staff_change.record(db, t, f"{_span(t, a.start_day, a.end_day)}: {a.role.replace('_', ' ')} {a.vendor.name} removed", ctx.user.id)
    db.delete(a)
    db.commit()


# --- participants -------------------------------------------------------------------

@router.get("/participants", response_model=list[ParticipantOut])
def list_participants(tour_id: str, ctx: WorkspaceCtx = Depends(office), db: Session = Depends(get_db)):
    return _tour(db, ctx, tour_id).participants


def next_position(db: Session, tour_id: str) -> int:
    return (db.scalar(select(func.max(Participant.position)).where(Participant.tour_id == tour_id)) or 0) + 1


@router.post("/participants", response_model=ParticipantOut, status_code=status.HTTP_201_CREATED)
def add_participant(tour_id: str, body: ParticipantIn, ctx: WorkspaceCtx = Depends(office),
                    db: Session = Depends(get_db)):
    t = _tour(db, ctx, tour_id)
    p = Participant(tour_id=t.id, position=next_position(db, t.id), **body.model_dump(exclude_none=True))
    db.add(p)
    db.commit()
    return p


def _participant(db: Session, t: Tour, pid: str) -> Participant:
    p = db.scalar(select(Participant).where(Participant.id == pid, Participant.tour_id == t.id))
    if not p:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Participant not found")
    return p


@router.patch("/participants/{participant_id}", response_model=ParticipantOut)
def update_participant(tour_id: str, participant_id: str, body: ParticipantUpdate,
                       ctx: WorkspaceCtx = Depends(office), db: Session = Depends(get_db)):
    p = _participant(db, _tour(db, ctx, tour_id), participant_id)
    for k, v in body.model_dump(exclude_unset=True).items():
        # Dates may be cleared with null; text fields only change to a value.
        if v is not None or k in ("birth_date", "passport_expiry"):
            setattr(p, k, v)
    db.commit()
    return p


@router.delete("/participants/{participant_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_participant(tour_id: str, participant_id: str, ctx: WorkspaceCtx = Depends(office),
                       db: Session = Depends(get_db)):
    db.delete(_participant(db, _tour(db, ctx, tour_id), participant_id))
    db.commit()


@router.delete("/participants", status_code=status.HTTP_204_NO_CONTENT)
def clear_participants(tour_id: str, ctx: WorkspaceCtx = Depends(office), db: Session = Depends(get_db)):
    t = _tour(db, ctx, tour_id)
    if t.closed_at is not None:       # a bulk delete skips the per-row lock: check here
        raise TourLocked(t.code)
    db.execute(delete(Participant).where(Participant.tour_id == t.id))
    db.commit()


# --- rooming -------------------------------------------------------------------------

def rooming(t: Tour) -> RoomingSummary:
    """Rooms needed: one room per room_group label; people without a label are
    packed by their room type's capacity (two DBL people = one DBL room)."""
    rooms: Counter[str] = Counter()
    ungrouped: Counter[str] = Counter()
    groups: dict[str, list[Participant]] = defaultdict(list)
    unassigned = 0
    for p in t.participants:
        if p.room_group.strip():
            groups[p.room_group.strip().lower()].append(p)
        elif p.room_type:
            ungrouped[p.room_type] += 1
        else:
            unassigned += 1
    for members in groups.values():
        types = Counter(m.room_type for m in members if m.room_type)
        rooms[types.most_common(1)[0][0] if types else _type_for_size(len(members))] += 1
    for rtype, count in ungrouped.items():
        rooms[rtype] += math.ceil(count / ROOM_CAPACITY[rtype])
    ordered = {r: rooms[r] for r in ("DBL", "TWN", "SGL", "TRP", "QUAD") if rooms[r]}
    limit = t.end_date + PASSPORT_MARGIN
    warnings = sum(1 for p in t.participants if p.passport_expiry and p.passport_expiry < limit)
    return RoomingSummary(rooms=ordered, text=" + ".join(f"{n} {r}" for r, n in ordered.items()),
                          unassigned=unassigned, passport_warnings=warnings)


def _type_for_size(n: int) -> str:
    return {1: "SGL", 2: "DBL", 3: "TRP"}.get(n, "QUAD")


@router.get("/rooming", response_model=RoomingSummary)
def rooming_summary(tour_id: str, ctx: WorkspaceCtx = Depends(office), db: Session = Depends(get_db)):
    return rooming(_tour(db, ctx, tour_id))


@router.get("/rooming-list.xlsx")
def rooming_list(tour_id: str, ctx: WorkspaceCtx = Depends(office), db: Session = Depends(get_db)):
    """The list hotels and border posts ask for, grouped by room."""
    t = _tour(db, ctx, tour_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rooming list"
    ws.append([f"{t.code}  ·  {t.client.name if t.client else ''}  ·  "
               f"{t.start_date:%d.%m.%Y} – {t.end_date:%d.%m.%Y}  ·  {len(t.participants)} pax  ·  {rooming(t).text}"])
    ws["A1"].font = Font(bold=True, size=12)
    header = ["#", "Room", "Room type", "Full name", "Gender", "Date of birth", "Nationality",
              "Passport", "Passport expiry", "Phone", "Notes"]
    ws.append([])
    ws.append(header)
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(vertical="center")
    # Roommates next to each other; unlabelled people keep list order at the end.
    people = sorted(t.participants, key=lambda p: (p.room_group.strip() == "", p.room_group.strip().lower(), p.position))
    for i, p in enumerate(people, 1):
        ws.append([i, p.room_group, p.room_type, p.full_name, p.gender, p.birth_date, p.nationality,
                   p.passport_number, p.passport_expiry, p.phone, p.notes])
    for row in ws.iter_rows(min_row=4):
        for c in (row[5], row[8]):
            c.number_format = "DD.MM.YYYY"
    for col, width in zip("ABCDEFGHIJK", (5, 10, 10, 30, 8, 14, 16, 16, 16, 18, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A4"
    buf = io.BytesIO()
    defuse(wb)
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{_safe(t.code)}-rooming-list.xlsx"'})


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]", "_", name) or "tour"

